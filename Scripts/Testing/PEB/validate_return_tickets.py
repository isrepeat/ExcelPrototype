"""Check the saved Movement return-ticket contract without Office or VBA."""

from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
from shutil import copyfile
import re

from openpyxl import Workbook, load_workbook
from config_dependencies import load_config


ROOT = Path(__file__).resolve().parents[3]
VBA = ROOT / "PROJECTS_EXCEL/vba/PersonalEventBuilder"
MODES = {
    "FromVacationForCuring", "FromAnnualMainVacation",
    "FromVacationForFamily", "FromBusinessTrip",
}


def config(name):
    values = load_config(ROOT / "PROJECTS_EXCEL/config/PersonalEventBuilder" / name)
    return {key.removeprefix("PersonalEventBuilder::"): value for key, value in values.items()}


def resolve(rows, orders, ipn, fio, year):
    """Reference the legacy reverse physical-row and current-year order rules."""
    temporary = ipn.strip().casefold() == "тп-іпн"
    pair = next((row[2:] for row in reversed(rows)
                 if (str(row[1]).strip().casefold() == fio.strip().casefold() if temporary
                     else str(row[0]).strip() == ipn.strip())
                 and row[2] and row[3]), None)
    if pair is None:
        return "", ""
    ticket, number = pair
    matches = {day for order, day in orders if int(order) == int(number) and day.year == year}
    if len(matches) > 1:
        raise ValueError("Conflicting order dates")
    if not matches:
        return "", ""
    return ticket, next(iter(matches)).strftime("%d.%m.%Y")


def main():
    for name in ("wsConfig.txt", "wsConfig.test.txt"):
        cfg = config(name)
        assert cfg["lookup.ReturnTickets.document"] == "В/к"
        assert "Ticket" not in cfg["extensions.types"].split("|")
        assert not any(key.startswith("extensions.type.") and key.endswith(".binding") for key in cfg)
        for form in cfg["export.ticket.form"].split("|"):
            fields = {cfg[f"forms.{form}.fields.{index}.name"]:
                      cfg.get(f"forms.{form}.fields.{index}.record", "")
                      for index in range(1, int(cfg[f"forms.{form}.fields.count"]) + 1)}
            assert not {"TicketNumber", "TicketDate"} & fields.keys()
            assert fields["Document"] == "Event.Document"
            assert fields["DocumentNumber"] == "Absence.TicketNumber"
            assert fields["DocumentDate"] == "Absence.TicketDate"
        assert set(cfg["lookup.ReturnTickets.modes"].split("|")) == MODES
        assert cfg["lookup.ReturnTickets.ticketHeader"] == "Супровідний документ"
        assert cfg["lookup.ReturnTickets.orderHeader"] == "Вибуття.Наказ"
        assert cfg["lookup.ReturnTickets.source.range"] == cfg["validation.Movement.source.range"]
        assert cfg["lookup.ReturnTickets.temporaryIpn"] == "ТП-ІПН"
        if name.endswith("test.txt"):
            assert "tests/PersonalEventBuilder/Dependencies/test_РУХ.xlsx" in Path(cfg["lookup.ReturnTickets.source.path"]).as_posix()
        else:
            assert Path(cfg["lookup.ReturnTickets.source.path"]).name == "РУХ_last.xlsx"

    provider = (VBA / "obj_PEB_ReturnTickets.cls.vba").read_text(encoding="utf-8-sig")
    enricher = (VBA / "obj_PEB_CandidateEnricher.cls.vba").read_text(encoding="utf-8-sig")
    controller = (VBA / "obj_PEB_PgMainController.cls.vba").read_text(encoding="utf-8-sig")
    exporter = (VBA / "obj_PEB_ExportService.cls.vba").read_text(encoding="utf-8-sig")
    assert 'm_loadingSnapshot.TryQuery(source, query, table, diagnostic)' in provider
    assert 'CopyFile' not in provider and 'GetFile' not in provider
    assert "For row = table.RowCount To 1 Step -1" in provider
    assert 'VBA.Len(ticket) > 0 And VBA.Len(orderNumber) > 0' in provider
    assert 'm_byIpn.Add ipn, VBA.Array(ticket, orderNumber, destination, arrivalValue)' in provider
    assert 'm_byName.Add name, VBA.Array(ticket, orderNumber, destination, arrivalValue)' in provider
    assert 'fields("Absence.PlannedArrivalDate") = Format$(CDate(pair(3)), "dd.mm.yyyy")' in provider
    assert 'fields("Event.IncomingDate") = VBA.Format$(parsedOrderDate, "dd.mm.yyyy")' in enricher
    assert 'fields("Absence.Destination") = pair(2)' in provider
    assert config('wsConfig.txt')['lookup.ReturnTickets.destinationHeader'] == 'Куди'
    assert config('wsConfig.test.txt')['lookup.ReturnTickets.destinationHeader'] == 'Куди'
    assert 'VBA.Replace$(ex_PEB_Export.fn_Config(prefix & ".orderHeader"), ".", "#")' in provider
    handler = provider.split("Public Function TryEnrich(", 1)[1].split("End Function", 1)[0].split("EH:", 1)[1]
    assert handler.index('errorText = VBA.Err.Description') < handler.index('ex_PEB_Export.fn_Config(')
    assert 'RETURN_TICKETS_FAILED | Stage=' in handler
    ado_headers = ["ІПН", "ПІБ", "Супровідний документ", "Вибуття#Наказ"]
    assert config("wsConfig.txt")["lookup.ReturnTickets.orderHeader"].replace(".", "#") in ado_headers
    assert 'VBA.Year(orderDate) = m_orderYear' in provider
    assert 'snapshot.Prepared("ReturnTickets")' in enricher
    assert 'Stamp' not in provider and 'PrepareSnapshot' in provider
    assert not re.search(r"Workbooks|SaveCopyAs|AcceptedOrder", provider)
    assert 'If Not returnMode And record.TryGetValue("Absence.TicketSequence"' in enricher
    assert 'm_returnTickets.TryEnrich(output, diagnostic)' in enricher
    assert 'snapshot.Prepared("ReturnTickets")' in enricher
    mode_loader = controller.split("Private Function private_LoadMode(", 1)[1].split("End Function", 1)[0]
    assert 'enrichedRecords' in mode_loader and 'm_candidateEnricher.TryEnrich(' in mode_loader
    mode_switch = controller.split("Private Function private_SwitchMode(", 1)[1].split("End Function", 1)[0]
    assert mode_switch.index('updates("Form.TicketNumber")') < mode_switch.index('private_LoadMode(modeId')
    assert 'Me.SelectPersonnelHandler(record)' in mode_switch
    ticket_export = exporter.split("Private Sub private_PrepareTicket(", 1)[1].split("End Sub", 1)[0]
    assert ticket_export.index('lookup.ReturnTickets.modes') < ticket_export.index('pattern.Pattern')
    export_helpers = (VBA / "ex_PEB_Export.vba").read_text(encoding="utf-8-sig")
    assert 'values("TicketNumber") = values("DocumentNumber")' in export_helpers
    assert 'values("TicketDate") = values("DocumentDate")' in export_helpers
    assert 'values("DocumentNumber") = legacyNumber' in export_helpers
    readonly = controller.split("Private Function private_UpdateReadOnly(", 1)[1].split("End Function", 1)[0]
    assert 'updates("Form.DocumentReadOnly") = True' in readonly
    assert 'updates("Form.Document") = ex_PEB_Export.fn_FixedDocument(modeId)' in readonly
    assert '.binding' not in controller

    year = 2026
    rows = [
        ("123", "PERSON A", "2026/7/1", "7"),
        ("123", "PERSON A", "2026/8/2", "8"),
        ("123", "PERSON A", "", "9"),
        ("123", "PERSON A", "partial ticket", ""),
        ("456", "PERSON B", "legacy literal ticket", "7"),
    ]
    orders = [(7, date(year - 1, 1, 1)), (7, date(year, 2, 3)), (8, date(year, 4, 5))]
    assert resolve(rows, orders, "123", "PERSON A", year) == ("2026/8/2", "05.04.2026")
    assert resolve(rows, orders, "ТП-ІПН", "person b", year) == ("legacy literal ticket", "03.02.2026")
    assert resolve(rows, orders, "missing", "PERSON A", year) == ("", "")
    assert resolve(rows, orders[:1], "123", "PERSON A", year) == ("", "")
    assert resolve(rows, orders + [(8, date(year, 4, 5))], "123", "PERSON A", year)[0] == "2026/8/2"
    try:
        resolve(rows, orders + [(8, date(year, 4, 6))], "123", "PERSON A", year)
    except ValueError:
        pass
    else:
        raise AssertionError("Conflicting departure-order dates must stop enrichment")

    with TemporaryDirectory() as folder:
        original = Path(folder) / "Movement.xlsx"
        snapshot = Path(folder) / "snapshot.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["IPN", "FIO", "Ticket", "Order"])
        for row in rows:
            sheet.append(row)
        workbook.save(original)
        sheet.append(["123", "PERSON A", "unsaved ticket", "7"])
        copyfile(original, snapshot)
        saved = load_workbook(snapshot, read_only=True, data_only=True)
        assert resolve(list(saved.active.values)[1:], orders, "123", "PERSON A", year)[0] == "2026/8/2"
        saved.close()
        workbook.save(original)
        copyfile(original, snapshot)
        saved = load_workbook(snapshot, read_only=True, data_only=True)
        assert resolve(list(saved.active.values)[1:], orders, "123", "PERSON A", year) == ("unsaved ticket", "03.02.2026")
        saved.close()
        workbook.close()

    print("PASS: PROD/TEST routes, snapshot wiring, mode changes, paired history, temporary IPN, order dates, saved/unsaved source contract. VBA was not executed.")


if __name__ == "__main__":
    main()