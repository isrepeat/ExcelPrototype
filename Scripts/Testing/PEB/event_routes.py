"""Read and validate the shared business-event registry for offline fixture tools."""

from pathlib import Path
from xml.etree import ElementTree as ET
import re

from movement_templates import load_movement
from movement_plans import load_native_plans

ROOT = Path(__file__).resolve().parents[3]
WORKBOOK = ROOT / "PROJECTS_EXCEL/workbook/PersonalEventBuilder"
NS = {"p": "urn:excelprototype:profiles"}


def parse_config(path):
    result = {}
    for number, line in enumerate(Path(path).read_text(encoding="utf-8-sig").splitlines(), 1):
        if not line:
            continue
        parts = line.split("\t")
        if len(parts) != 3:
            raise ValueError(f"Invalid config row: {path}:{number}")
        key = parts[1].split("::", 1)[-1]
        if key in result:
            raise ValueError(f"Duplicate config key: {key}")
        result[key] = parts[2]
    return result


def validate_word_template(template):
    def require(condition, message):
        if not condition:
            raise ValueError("Word template XML: " + message)
    def attributes(node, allowed, required=()):
        require(set(node.attrib) <= set(allowed), "Unknown attribute")
        require(all(node.get(key, "").strip() for key in required), "Missing attribute")
    def text(node):
        attributes(node, ())
        require(not list(node), "Text must contain DSL, not XML elements")
    def text_block(node):
        require(len(node) == 1 and node[0].tag == f"{{{NS['p']}}}text", "Expected one text element")
        text(node[0])
    def group(node):
        attributes(node, ("by",), ("by",))
        require(all(key.strip() for key in node.get("by").split(";")), "Empty grouping key")
        require(all(child.tag in {f"{{{NS['p']}}}{tag}" for tag in ("aggregate", "single", "multiple")} for child in node), "Unknown group element")
        require(len(node.findall("p:single", NS)) == len(node.findall("p:multiple", NS)) == 1, "Expected single and multiple")
        single = node.find("p:single", NS)
        attributes(single, ())
        text_block(single)
        multiple = node.find("p:multiple", NS)
        attributes(multiple, ())
        require(len(multiple.findall("p:item", NS)) + len(multiple.findall("p:group", NS)) == 1, "Expected one item or nested group")
        for child in multiple:
            if child.tag == f"{{{NS['p']}}}group":
                group(child)
            else:
                require(child.tag in {f"{{{NS['p']}}}{tag}" for tag in ("header", "item", "footer")}, "Unknown multiple element")
                require(sum(other.tag == child.tag for other in multiple) == 1, "Duplicate group block")
                common = ("common",) if child.tag.endswith(("}header", "}footer")) else ()
                attributes(child, common, common)
                text_block(child)
        names = set()
        for aggregate in node.findall("p:aggregate", NS):
            keys = ("name", "field", "operation", "pattern", "separator", "rangeSeparator")
            attributes(aggregate, keys, keys)
            require(not list(aggregate), "Aggregate cannot contain elements")
            require(aggregate.get("name") not in names, "Duplicate aggregate")
            names.add(aggregate.get("name"))
            require(aggregate.get("operation") == "numberRanges", "Unknown aggregate operation")
    attributes(template, ("name", "hash", "identityBy", "orderBy"), ("name",))
    bodies = template.findall("p:text", NS) + template.findall("p:group", NS)
    require(len(bodies) == 1, "Expected exactly one text or group body")
    body_seen = False
    for child in template:
        if child.tag == f"{{{NS['p']}}}derived":
            require(not body_seen, "Derived fields must precede the body")
            attributes(child, ("name",), ("name",))
            text_block(child)
        else:
            require(child in bodies, "Unknown template element")
            body_seen = True
    if bodies[0].tag == f"{{{NS['p']}}}text":
        require(not {"identityBy", "orderBy"} & template.attrib.keys(), "State keys require a group")
        text(bodies[0])
    else:
        require(bool(template.get("identityBy")) and bool(template.get("orderBy")), "Missing state keys")
        group(bodies[0])


def load_registry(settings, xml_text=None, movement_text=None, word_text=None):
    path = WORKBOOK / settings["events.path"]
    text = xml_text if xml_text is not None else path.read_text(encoding="utf-8-sig")
    if "<!DOCTYPE" in text.upper() or "<!ENTITY" in text.upper():
        raise ValueError("Event registry DTD/entities are prohibited")
    root = ET.fromstring(text)
    def require(condition, message):
        if not condition:
            raise ValueError(f"Event routes XML: {message}")
    def attributes(node, allowed, required):
        require(set(node.attrib) <= set(allowed.split("|")), f"Unknown attribute on {node.tag}")
        for key in required.split("|"):
            require(bool(node.get(key, "").strip()), f"Missing {node.tag}/@{key}")
    def index(nodes, key):
        result = {}
        for node in nodes:
            identifier = node.get(key, "")
            require(re.fullmatch(r"[A-Za-z][A-Za-z0-9_]*", identifier), f"Invalid identifier: {identifier}")
            require(identifier not in result, f"Duplicate identifier: {identifier}")
            result[identifier] = node
        return result
    require(root.tag == f"{{{NS['p']}}}eventRoutes", "Expected eventRoutes root")
    attributes(root, "version|bindingField", "version|bindingField")
    require(root.get("version") == "1", "Unsupported version")
    fields = set(settings["export.fields"].split("|"))
    require(root.get("bindingField") in fields, "Unknown binding field")
    for tag in ("events",):
        require(len(root.findall(f"p:{tag}", NS)) == 1, f"Expected one {tag}")
    require({child.tag for child in root} == {f"{{{NS['p']}}}{tag}" for tag in ("events",)}, "Unknown root element")
    events = index(root.findall("p:events/p:event", NS), "id")
    require(events, "Empty event registry")
    has_movement = any(event.find("p:movement", NS) is not None for event in events.values())
    movement_registry = load_movement(settings, WORKBOOK, movement_text) if has_movement else {
        "statuses": {}, "values": {}, "templates": {},
    }
    if has_movement:
        movement_registry.update(load_native_plans(settings))
    statuses = {identifier: ET.Element(f"{{{NS['p']}}}status", id=identifier, name=name)
                for identifier, name in movement_registry["statuses"].items()}
    for node in statuses.values():
        attributes(node, "id|name", "id|name")
    require(len({node.get('name') for node in statuses.values()}) == len(statuses), "Duplicate Movement status name")
    templates = {}
    if any(event.find("p:word", NS) is not None for event in events.values()):
        text = word_text if word_text is not None else (WORKBOOK / settings["export.Word.templates.path"]).read_text(encoding="utf-8-sig")
        require("<!DOCTYPE" not in text.upper() and "<!ENTITY" not in text.upper(), "Word DTD/entities are prohibited")
        word_root = ET.fromstring(text)
        require(word_root.tag == f"{{{NS['p']}}}wordResultTemplates" and word_root.get("version") == "2", "Invalid Word root/version")
        require(all(child.tag in {f"{{{NS['p']}}}{name}" for name in ("documentFilepath", "sharedTemplates", "template")} for child in word_root),
                "Unknown Word template section")
        templates = index(word_root.findall("p:template", NS), "name")
        for template in templates.values():
            validate_word_template(template)
    for identifier, event in events.items():
        attributes(event, "id|caption|label|direction|form|visible", "id|caption|label|direction|form|visible")
        require(event.get("direction") in ("Arrival", "Departure", "StatusChange"), f"Invalid direction: {identifier}")
        require(event.get("form") in ("Leave", "Treatment"), f"Invalid form: {identifier}")
        require(event.get("visible") in ("true", "false"), f"Invalid visibility: {identifier}")
        tags = [child.tag for child in event]
        allowed = {f"{{{NS['p']}}}movement", f"{{{NS['p']}}}word"}
        require(tags and len(tags) == len(set(tags)) and set(tags) <= allowed,
                f"Expected one or both independent exporter routes: {identifier}")
        movement = event.find("p:movement", NS)
        word = event.find("p:word", NS)
        if movement is not None:
            attributes(movement, "template", "template")
            require(not list(movement), f"Unexpected Movement route children: {identifier}")
            name = movement.get("template")
            require(name in movement_registry["templates"], f"Unknown Movement template: {identifier}/{name}")
            require(identifier in movement_registry["events"], f"Unknown native Movement event: {identifier}")
            plan = movement_registry["events"][identifier]
            prefixes = ["profiles." + plan["profile"] + ".fields"]
            values = movement_registry["values"]
            if plan["operation"] == "Transition":
                prefixes.append("profiles." + values["profiles." + plan["profile"] + ".closeProfile"] + ".fields")
            allowed_columns, required_columns = set(), set()
            for prefix in prefixes:
                for index in range(1, int(values[prefix + ".count"]) + 1):
                    header = values[f"{prefix}.{index}.header"].casefold()
                    allowed_columns.add(header)
                    if values[f"{prefix}.{index}.field"] == "MovementBasis":
                        required_columns.add(header)
            provided_columns = {column.casefold() for column in movement_registry["columns"][name]}
            require(provided_columns <= allowed_columns, f"Text column is not writable: {identifier}")
            require(required_columns <= provided_columns, f"Missing required text column: {identifier}")
            movement.attrib.update(plan)
        if word is not None:
            attributes(word, "template", "template")
            require(not list(word), f"Unexpected Word route children: {identifier}")
            require(word.get("template") in templates, f"Unknown Word template: {identifier}")
            profile = "export.Word." + word.get("template")
            for suffix in ("required", "previewRequired"):
                require(bool(settings.get(profile + "." + suffix)), f"Unknown Word preparation settings: {identifier}")
            require(settings.get(profile + ".lookups.count", "").isdigit(), f"Invalid Word preparation settings: {identifier}")
    return {"events": events, "statuses": statuses, "bindingField": root.get("bindingField"), "movement": movement_registry}


def derived_settings(settings):
    """Expose an in-memory compatibility view; never persist derived keys to wsConfig."""
    registry = load_registry(settings)
    result = dict(settings)
    result.update({"export.Movement." + key: value for key, value in registry["movement"]["values"].items()})
    events = registry["events"]
    result["events.ids"] = result["test.registry.ModeId.values"] = "|".join(events)
    for identifier, event in events.items():
        result[f"test.registry.ModeId.{identifier}.runtimeMode"] = identifier
        for key in ("caption", "form", "direction"):
            result[f"modes.{identifier}.{key}"] = event.get(key)
        result[f"export.events.{identifier}"] = event.get("label")
        movement = event.find("p:movement", NS)
        word = event.find("p:word", NS)
        result[f"events.{identifier}.word.policy"] = "Render" if word is not None else "Omit"
        if word is not None:
            for key, value in word.attrib.items():
                result[f"events.{identifier}.word.{key}"] = value
        if movement is not None:
            for key, value in movement.attrib.items():
                result[f"events.{identifier}.movement.{key}"] = value
            result[f"export.Movement.{identifier}.profile"] = movement.get("profile")
            result[f"export.Movement.{identifier}.required"] = movement.get("required")
            result[f"export.Movement.{identifier}.destinationField"] = movement.get("destinationField")
    return result


if __name__ == "__main__":
    import json
    import sys
    settings = parse_config(sys.argv[-1])
    if "--orders" in sys.argv:
        import datetime
        import openpyxl
        from openpyxl.utils.cell import range_boundaries
        from openpyxl.utils.datetime import to_excel
        book = openpyxl.load_workbook(WORKBOOK / settings["orders.source.path"], data_only=True)
        sheet = book[settings["orders.source.sheet"]]
        left, top, right, bottom = range_boundaries(settings["orders.source.range"])
        headers = [sheet.cell(top, column).value for column in range(left, right + 1)]
        dates = left + headers.index(settings["orders.dateHeader"])
        numbers = left + headers.index(settings["orders.numberHeader"])
        result = {}
        for row in range(top + 1, min(bottom, sheet.max_row) + 1):
            date, number = sheet.cell(row, dates).value, sheet.cell(row, numbers).value
            if date is None and number is None:
                continue
            if isinstance(date, str):
                date = datetime.datetime.strptime(date, "%d.%m.%Y")
            serial = int(to_excel(date, book.epoch)) if isinstance(date, (datetime.datetime, datetime.date)) else int(date)
            values = result.setdefault(str(serial), [])
            if str(number) not in values:
                values.append(str(number))
        print(json.dumps(result, ensure_ascii=False))
    else:
        print(json.dumps(derived_settings(settings), ensure_ascii=False))