"""Migrate grouping fixtures to form inputs; preserve unrelated XLSX package parts."""

import argparse
import hashlib
import json
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl
from openpyxl.utils.cell import range_boundaries, get_column_letter

from event_routes import ROOT, WORKBOOK, parse_config
from update_grouping_visuals import apply, NS
from word_grouping_fixtures import load_personnel


def plan():
    settings = parse_config(ROOT / "PROJECTS_EXCEL/config/PersonalEventBuilder/wsConfig.test.txt")
    path = (WORKBOOK / settings["test.scenarios.path"]).resolve()
    book = openpyxl.load_workbook(path)
    people = load_personnel(settings)
    changes = []
    tables = []
    sheets = {}
    aliases = {"ModeId": "EventId", "AcceptedOrderDate": "OrderDateShort", "EventDate": "DateFromShort",
               "TicketNumber": "VacationTicketNo", "TicketDate": "VacationTicketDateShort",
               "IncomingNumber": "IncomingNo", "IncomingDate": "IncomingDateShort"}
    replacements = {"ТЕСТЕНКА Андрія": people["ipn_0001"]["PersonAccusative"],
                    "ПРИКЛАДЕНКА Богдана": people["ipn_0002"]["PersonAccusative"]}
    for kind in ("groupingSteps", "groupingExpected"):
        sheet = book[settings[f"test.scenarios.{kind}.sheet"]]
        table = sheet.tables[settings[f"test.scenarios.{kind}.table"]]
        left, top, right, bottom = range_boundaries(table.ref)
        old_headers = [sheet.cell(top, column).value for column in range(left, right + 1)]
        headers = settings[f"test.scenarios.{kind}.fields"].split("|")
        sheets[sheet.title] = []
        for row in range(top, bottom + 1):
            record = dict(zip(old_headers, [sheet.cell(row, column).value for column in range(left, right + 1)]))
            if row == top:
                values = headers
            elif kind == "groupingSteps":
                if "EventId" not in record:
                    values = [record[field] for field in headers]
                else:
                    record["ReporterId"] = "" if record["ReportPerson"] in (None, "", "сам") else "ipn_0005"
                    values = [record.get(aliases.get(field, field)) for field in headers]
            else:
                text = record["Text"]
                for original, replacement in replacements.items():
                    text = text.replace(original, replacement)
                if record["ScenarioId"] == "G02-second-person-rebuilds" and text in replacements.values():
                    person_id = "ipn_0001" if text == replacements["ТЕСТЕНКА Андрія"] else "ipn_0002"
                    person = people[person_id]
                    position = person["PositionGenitive"]
                    text = f'{person["RankGenitive"]} {text}, {person_id}, {position[:1].lower() + position[1:]}'
                record["Text"] = text
                values = [record[field] for field in headers]
            sheets[sheet.title].append(["" if value is None else value for value in values])
            for offset in range(right - left + 1):
                cell = sheet.cell(row, left + offset)
                value = values[offset] if offset < len(values) else ""
                value = "" if value is None else value
                if (cell.value if cell.value is not None else "") != value:
                    assert cell.data_type != "f"
                    changes.append({"sheet": sheet.title, "cell": cell.coordinate, "old": cell.value, "value": value})
        if old_headers != headers:
            tables.append({"name": table.name, "headers": headers,
                           "ref": f'{get_column_letter(left)}{top}:{get_column_letter(left + len(headers) - 1)}{bottom}'})
    book.close()
    with ZipFile(path) as package:
        for name in package.namelist():
            if name.startswith("xl/tables/") and name.endswith(".xml"):
                table = ET.fromstring(package.read(name))
                for change in tables:
                    if table.get("name") == change["name"]:
                        change["part"] = name
    assert all("part" in change for change in tables)
    return {"path": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "changes": changes, "tables": tables, "sheets": sheets}


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--candidate")
    args = parser.parse_args()
    edits = plan()
    if args.candidate:
        apply(args.candidate, edits)
    else:
        print(json.dumps(edits, ensure_ascii=True))