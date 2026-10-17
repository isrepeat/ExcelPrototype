"""Check XML routes, malformed references and fixture preservation without Office."""

from copy import deepcopy
from pathlib import Path
from xml.etree import ElementTree as ET
import json
import subprocess
import zipfile

import openpyxl
from openpyxl.utils.datetime import to_excel

from event_routes import ROOT, NS, WORKBOOK, parse_config, load_registry
from xlsx_expectations_bridge import validate_fixture_layout


def main():
    config = ROOT / "PROJECTS_EXCEL/config/PersonalEventBuilder"
    test = parse_config(config / "wsConfig.test.txt")
    prod = parse_config(config / "wsConfig.txt")
    assert test["events.path"] == prod["events.path"]
    for settings in (test, prod):
        assert not any(key.startswith(("events.", "export.events.", "test.registry.ModeId.")) and key not in ("events.path", "test.registry.ModeId.rangeName") for key in settings)
        assert not any(key.startswith("modes.") and key != "modes.default" for key in settings)
        registry = load_registry(settings)
        assert settings["modes.default"] in registry["events"]
    events = registry["events"]
    special = events["MoveFromOutpatientVlkToVacationForCuring"]
    assert special[0].get("operation") == "Transition"
    assert special[0].get("sourceStatus") == "OutpatientVlk"
    assert special[0].get("targetStatus") == "MedicalLeave"
    assert special[1].get("template") == "ToVacationForCuring"
    for identifier in ("FromOutpatientExaminationOrVlk", "ToOutpatientExaminationOrVlk"):
        assert events[identifier][1].get("policy") == "Omit"
        assert events[identifier][1].get("template") is None
    root = ET.parse(WORKBOOK / test["events.path"]).getroot()
    rejected = 0
    for section, attribute, value in (
        ("movement", "sourceStatus", "Missing"),
        ("movement", "targetStatus", "Missing"),
        ("movement", "operation", "Missing"),
        ("movement", "profile", "Missing"),
        ("movement", "required", "UnknownInput"),
        ("word", "profile", "Missing"),
        ("word", "template", "Missing"),
        ("word", "anchor", "Missing"),
        ("word", "policy", "Missing"),
    ):
        altered = deepcopy(root)
        node = altered.find("p:events/p:event", NS).find("p:" + section, NS)
        node.set(attribute, value)
        try:
            load_registry(test, ET.tostring(altered, encoding="unicode"))
        except ValueError:
            rejected += 1
        else:
            raise AssertionError(f"Invalid XML reference accepted: {section}/{attribute}")
    for selector in ("p:events", "p:movementStatuses", "p:wordAnchors"):
        altered = deepcopy(root)
        container = altered.find(selector, NS)
        container.append(deepcopy(container[0]))
        try:
            load_registry(test, ET.tostring(altered, encoding="unicode"))
        except ValueError:
            rejected += 1
        else:
            raise AssertionError(f"Duplicate XML identifier accepted: {selector}")
    fixture = WORKBOOK / test["test.scenarios.path"]
    with zipfile.ZipFile(fixture) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    validate_fixture_layout(parts)
    assert "xl/calcChain.xml" not in parts
    broken = dict(parts)
    changed = False
    for name, content in parts.items():
        if name.startswith("xl/tables/") and name.endswith(".xml"):
            node = ET.fromstring(content)
            if node.get("name") == test["test.scenarios.expectedMovement.table"]:
                node.set("ref", "A1:E44")
                broken[name] = ET.tostring(node)
                changed = True
    assert changed
    try:
        validate_fixture_layout(broken)
    except ValueError:
        rejected += 1
    else:
        raise AssertionError("Caption collision accepted")
    book = openpyxl.load_workbook(fixture)
    orders = json.loads(subprocess.check_output([__import__("sys").executable, "-X", "utf8", str(Path(__file__).with_name("event_routes.py")), "--orders", str(config / "wsConfig.test.txt")], text=True, encoding="utf-8"))
    rows = {}
    for name in ("tbTestHistory", "tbExpectedMovementRows", "tbTestSteps"):
        sheet = next(sheet for sheet in book if name in sheet.tables)
        cells = list(sheet[sheet.tables[name].ref])
        headers = [cell.value for cell in cells[0]]
        rows[name] = [dict(zip(headers, [cell.value for cell in row])) for row in cells[1:]]
    for row in rows["tbTestHistory"]:
        for date_key, order_key in (("Вибуття", "Вибуття.Наказ"), ("Прибуття", "Прибуття.Наказ")):
            if row[order_key]:
                serial = str(int(to_excel(row[date_key])))
                assert str(row[order_key]) in orders[serial], (row["ScenarioId"], date_key)
    duplicate = next(row for row in rows["tbTestHistory"] if row["ScenarioId"] == "S14")
    step = next(row for row in rows["tbTestSteps"] if row["ScenarioId"] == "S14")
    assert duplicate["Вибуття.Наказ"] == str(step["AcceptedOrderNumber"]) == "298"
    assert duplicate["Вибуття"] == step["EventDate"]
    for row in rows["tbExpectedMovementRows"]:
        if row["StepId"] in ("S16-M02", "S16-M03") and row["RecordId"] == "R01":
            assert str(row["Вибуття.Наказ"]) == "284"
            assert str(row["Прибуття.Наказ"]) == "298"
            assert row["Прибуття"].day == 14
    expected_sheet = book[test["test.scenarios.expectedRows.sheet"]]
    top = openpyxl.utils.cell.range_boundaries(expected_sheet.tables[test["test.scenarios.expectedRows.table"]].ref)[1]
    assert expected_sheet.cell(top - 1, 1).value == test["test.expected.rowsCaption"]
    print(f"PASS: {len(events)} XML events, {len(registry['statuses'])} unique statuses; {rejected} invalid routes/layouts rejected; historical orders, S14 and S16 verified.")


if __name__ == "__main__":
    main()