"""Проверка схем сценариев и контрактов исходного кода без запуска Office."""

import copy
import datetime
import re
import sys
from xml.etree import ElementTree as ET
from pathlib import Path

import openpyxl
from openpyxl.utils.cell import range_boundaries

ROOT = Path(__file__).resolve().parents[3]
FIXTURE = ROOT / "PROJECTS_EXCEL/tests/PersonalEventBuilder/TestScenarios.xlsx"


def config(filename):
    settings = {}
    seen = set()
    for number, line in enumerate(filename.read_text(encoding="utf-8-sig").splitlines(), 1):
        if not line.strip():
            continue
        parts = line.split("\t")
        assert len(parts) == 3, f"Invalid config row {filename}:{number}"
        assert parts[1] not in seen, f"Duplicate config key: {parts[1]}"
        seen.add(parts[1])
        settings[parts[1].split("::", 1)[-1]] = parts[2]
    return settings


def load_tables(book):
    result = {}
    for sheet in book:
        for table in sheet.tables.values():
            left, top, right, bottom = range_boundaries(table.ref)
            headers = [sheet.cell(top, column).value for column in range(left, right + 1)]
            assert headers == [column.name for column in table.tableColumns], f"Header schema differs: {table.name}"
            assert len(set(headers)) == len(headers), f"Duplicate header: {table.name}"
            rows = []
            for index in range(top + 1, bottom + 1):
                cells = [sheet.cell(index, column) for column in range(left, right + 1)]
                assert not any(cell.data_type == "e" for cell in cells), f"Excel error: {sheet.title}/{index}"
                row = {header: cell.value for header, cell in zip(headers, cells)}
                if any(value is not None for value in row.values()):
                    rows.append(row)
            result[table.name] = {"sheet": sheet.title, "headers": headers, "rows": rows}
    return result


def unique(rows, field):
    result = {}
    for row in rows:
        value = row[field]
        assert isinstance(value, str) and value.strip(), f"Missing identity: {field}"
        assert value not in result, f"Duplicate {field}: {value}"
        result[value] = row
    return result


def runtime_mode(step, settings):
    assert step["ModeId"] in settings["test.registry.ModeId.values"].split("|"), f"Unknown template ID: {step['ModeId']}"
    return settings[f"test.registry.ModeId.{step['ModeId']}.runtimeMode"]


def validate(tables, settings):
    cases = unique(tables["tbTestScenarios"]["rows"], "ScenarioId")
    steps = unique(tables["tbTestSteps"]["rows"], "StepId")
    movement = unique(tables["tbExpectedMovement"]["rows"], "StepId")
    word = unique(tables["tbExpectedWord"]["rows"], "StepId")
    snapshots = tables["tbExpectedMovementRows"]["rows"]
    for binding in settings["test.registry.bindings"].split("|"):
        table, field, kind = binding.split(":")
        allowed = settings[f"test.registry.{kind}.values"].split("|")
        assert allowed and len(allowed) == len(set(allowed)), f"Invalid registry: {kind}"
        assert all(row[field] in allowed for row in tables[table]["rows"]), f"Unknown reserved value: {table}/{field}"
    assert not {"ExpectedDecision", "ExpectedCode", "ExpectedOpenCount"}.intersection(tables["tbTestSteps"]["headers"])
    assert "PersonPosition" not in tables["tbTestSteps"]["headers"]
    assert set(tables["tbTestSteps"]["headers"]) == set(settings["test.scenarios.stepHeaders"].split("|"))
    assert "Цільова таблиця" not in tables["tbTestScenarios"]["headers"]
    for kind in ("cases", "steps", "history", "absence", "expectedMovement", "expectedRows", "expectedWord"):
        table = tables[settings[f"test.scenarios.{kind}.table"]]
        assert table["sheet"] == settings[f"test.scenarios.{kind}.sheet"], f"Wrong sheet for {kind}"
    for kind, key in (("movement", "tbExpectedMovement"), ("word", "tbExpectedWord")):
        assert tables[key]["headers"] == settings[f"test.expected.{kind}Headers"].split("|")
    native = [header for header in tables["tbTestHistory"]["headers"] if header not in ("ScenarioId", "RowId", "TargetTable")]
    helpers = settings["test.scenarios.formulaColumns"].split("|")
    assert tables["tbExpectedMovementRows"]["headers"] == ["StepId", "RecordId"] + [header for header in native if header not in helpers]
    for kind in ("results", "inputs", "records"):
        expected = settings[f"test.audit.{kind[:-1] if kind != 'records' else 'record'}Fields"].split("|")
        if kind == "records":
            expected += native
        assert tables[settings[f"test.audit.{kind}.table"]]["headers"] == expected, f"Archive schema differs: {kind}"
    for scenario_id, case in cases.items():
        assert case["RuleStatus"] in ("Defined", "PendingRule")
        orders = [row["StepOrder"] for row in steps.values() if row["ScenarioId"] == scenario_id]
        assert orders and sorted(orders) == list(range(1, len(orders) + 1)), f"Invalid StepOrder: {scenario_id}"
        assert any(row["ScenarioId"] == scenario_id for row in tables["tbTestAbsence"]["rows"]), f"Missing Absence fixture: {scenario_id}"
    for step_id, step in steps.items():
        assert step["ScenarioId"] in cases, f"Unknown scenario: {step_id}"
        assert step["PersonId"] == cases[step["ScenarioId"]]["ІПН"], f"Wrong participant: {step_id}"
        assert step["Action"] in ("ExportMovement", "PreviewWord"), f"Unknown action: {step_id}"
        expected = movement if step["Action"] == "ExportMovement" else word
        assert step_id in expected, f"Missing expectation: {step_id}"
        assert settings[f"export.events.{runtime_mode(step, settings)}"]
    for step_id, expectation in movement.items():
        assert step_id in steps and steps[step_id]["Action"] == "ExportMovement", f"Orphan Movement expectation: {step_id}"
        decision = expectation["ExpectedDecision"]
        assert decision in ("Allow", "Block", "NoChange", "Pending")
        assert expectation["ExpectedCode"], f"Missing code: {step_id}"
        state_check = "Snapshot" if decision == "Allow" else "Pending" if decision == "Pending" else "Unchanged"
        assert expectation["StateCheck"] == state_check, f"Incompatible state check: {step_id}"
        rows = [row for row in snapshots if row["StepId"] == step_id]
        if state_check == "Snapshot":
            assert rows, f"Missing snapshot: {step_id}"
            assert len({row["RecordId"] for row in rows}) == len(rows)
            keys = settings["test.expected.rowKeys"].split("|")
            assert len({tuple(row[field] for field in keys) for row in rows}) == len(rows), f"Duplicate snapshot identity: {step_id}"
            assert all(row["ІПН"] == steps[step_id]["PersonId"] for row in rows)
            assert sum(row["Прибуття"] is None for row in rows) == expectation["ExpectedOpenCount"], f"Wrong expected open count: {step_id}"
        else:
            assert not rows, f"Unexpected snapshot: {step_id}"
        if decision != "Pending":
            count = expectation["ExpectedOpenCount"]
            assert isinstance(count, (int, float)) and count >= 0 and count == int(count), f"Invalid open count: {step_id}"
    assert all(row["StepId"] in movement for row in snapshots), "Orphan expected row"
    for step_id, expectation in word.items():
        assert step_id in steps and steps[step_id]["Action"] == "PreviewWord", f"Orphan Word expectation: {step_id}"
        assert expectation["Outcome"] in ("Success", "Error", "Pending")
        if expectation["Outcome"] == "Success":
            text = expectation["ExpectedText"]
            assert isinstance(text, str) and 0 < len(text) <= 32767 and not expectation["ExpectedError"], f"Missing Word text: {step_id}"
        elif expectation["Outcome"] == "Error":
            assert expectation["ExpectedError"] and not expectation["ExpectedText"], f"Missing Word diagnostic: {step_id}"
    return len(cases), len(steps), len(snapshots), len(word)


def negative_checks(tables, settings):
    def reject(change):
        altered = copy.deepcopy(tables)
        change(altered)
        try:
            validate(altered, settings)
        except AssertionError:
            return
        raise AssertionError("Invalid fixture was accepted")

    reject(lambda data: data["tbTestSteps"]["rows"][1].update(StepId=data["tbTestSteps"]["rows"][0]["StepId"]))
    reject(lambda data: data["tbExpectedMovement"]["rows"][0].update(StepId="UNKNOWN"))
    reject(lambda data: data["tbExpectedMovement"]["rows"].pop(0))
    reject(lambda data: data["tbTestSteps"]["rows"][0].update(StepOrder=5))
    reject(lambda data: data["tbExpectedWord"]["rows"][0].update(ExpectedText=None))
    reject(lambda data: data["tbExpectedMovementRows"]["rows"].append(copy.deepcopy(data["tbExpectedMovementRows"]["rows"][0])))
    reject(lambda data: data["tbExpectedMovementRows"]["rows"][0].update(StepId="UNKNOWN"))
    reject(lambda data: data["tbExpectedMovementRows"]["rows"][0].update(ІПН="UNKNOWN"))
    reordered = copy.deepcopy(tables)
    for table in reordered.values():
        table["rows"].reverse()
    assert validate(reordered, settings) == validate(tables, settings), "Row sorting changed stable links"
    print("PASS: 8 invalid fixture cases rejected; sorting preserves links.")


def source_checks(settings):
    folder = ROOT / "PROJECTS_EXCEL/vba/PersonalEventBuilder"
    loader = (folder / "obj_PEB_TestScenarios.cls.vba").read_text(encoding="utf-8-sig")
    runner = (folder / "obj_PEB_TestRunner.cls.vba").read_text(encoding="utf-8-sig")
    catalog = (folder / "ex_PEB_Testing.vba").read_text(encoding="utf-8-sig")
    assert 'private_LoadExpectations allSteps' in loader
    assert 'private_Value(steps, stepRow, "expected")' not in loader
    assert 'private_Value(steps, stepRow, "code")' not in loader
    assert 'private_Value(steps, stepRow, "openCount")' not in loader
    assert 'private_ValidateRows(after, testStep, checkDiagnostic)' in runner
    assert 'private_ValidateWord(binding, testStep, report, values)' in runner
    assert 'VBA.StrComp(expected, actual, VBA.vbBinaryCompare)' in runner
    assert 'testStep("Order") = 1' in runner and 'testStep("Id") = 1' not in runner
    assert 'test.scenarios.headers.stepId' in catalog and 'TestPreviewOperation' in catalog
    service = (folder / "obj_PEB_ExportService.cls.vba").read_text(encoding="utf-8-sig")
    assert '"export.Movement." & mode & ".destinationField"' in service
    assert 'values("Destination") = ex_PEB_Export.fn_Required(values, VBA.CStr(key))' in service
    assert "HospitalShortName" in settings["export.Movement.ToHospital.required"].split("|")
    for mode in {runtime_mode(step, settings) for step in load_tables(openpyxl.load_workbook(FIXTURE))["tbTestSteps"]["rows"]}:
        assert settings[f"export.Movement.{mode}.destinationField"] == ("HospitalShortName" if mode in ("ToHospital", "FromHospital") else "Destination")
    for filename in ("obj_PEB_TestScenarios.cls.vba", "obj_PEB_TestRunner.cls.vba", "ex_PEB_Testing.vba", "obj_PEB_TestUi.cls.vba"):
        text = (folder / filename).read_text(encoding="utf-8-sig")
        procedures = re.findall(r"(?im)^(?:Public|Private|Friend) (?:Function|Sub) (\w+)", text)
        assert len(procedures) == len(set(procedures)), f"Duplicate procedure: {filename}"
        assert len(procedures) == len(re.findall(r"(?im)^End (?:Function|Sub)$", text)), f"Unbalanced procedures: {filename}"
        for name, body in re.findall(r"(?ms)^(?:Public|Private|Friend) (?:Function|Sub) (\w+)(.*?)(?=^End (?:Function|Sub)$)", text):
            declarations = re.findall(r"(?im)^\s*Dim (\w+)\b", body)
            assert len(declarations) == len(set(declarations)), f"Duplicate variable: {filename}/{name}"
        for key in re.findall(r'fn_Config\("([^"]+)"\)', text):
            assert key in settings, f"Missing configuration: {filename}/{key}"
    print("PASS: VBA source/configuration contracts and procedure structure.")


def word_context_checks(settings):
    root = ET.parse(ROOT / "PROJECTS_EXCEL/workbook/PersonalEventBuilder/PersonalEventBuilderWordResultTemplates.xml").getroot()
    ns = {"p": "urn:excelprototype:profiles"}
    text = "\n".join(root.itertext())
    roots = set(re.findall(r"(?<!\.)\[([A-Z][A-Za-z0-9_]*)\]", text))
    row_keys = set(re.findall(r"item\.\[([A-Z][A-Za-z0-9_]*)\]", text))
    provided = set(settings["export.fields"].split("|"))
    for group, name in [("lookups", "output"), ("aliases", "name"), ("constants", "name"), ("context.lookups", "output")]:
        for index in range(1, int(settings[f"export.Word.{group}.count"]) + 1):
            provided.add(settings[f"export.Word.{group}.{index}.{name}"])
    provided.update(settings["export.Word.context.optionalText"].split("|"))
    service = (ROOT / "PROJECTS_EXCEL/vba/PersonalEventBuilder/obj_PEB_ExportService.cls.vba").read_text(encoding="utf-8-sig")
    provided.update(re.findall(r'values\("([A-Za-z0-9_]+)"\)\s*=', service))
    assert roots <= provided, f"Missing Word context providers: {sorted(roots - provided)}"
    declared = set(settings["export.Word.context.requiredKeys"].split("|"))
    assert roots <= declared and declared <= provided
    for template in root.findall("p:template", ns):
        required = set(settings[f"export.Word.context.templates.{template.attrib['name']}.requiredKeys"].split("|"))
        assert required <= declared, f"Undeclared required Word fields: {template.attrib['name']}"
    meta_keys = set()
    for kind in settings["export.Word.context.collectionInputs"].split("|"):
        fields = set(settings[f"export.Word.context.{kind}.fields"].split("|"))
        required = set(settings[f"export.Word.context.{kind}.required"].split("|"))
        assert required <= fields
        meta_keys.update(fields)
    for index in range(1, int(settings["export.Word.meta.lookups.count"]) + 1):
        meta_keys.add(settings[f"export.Word.meta.lookups.{index}.output"])
    assert row_keys <= meta_keys, f"Missing metadata row providers: {row_keys - meta_keys}"
    xml_keys = set(settings["export.Word.context.xmlKeys"].split("|"))
    assert roots | row_keys <= xml_keys
    parser = (ROOT / "PROJECTS_EXCEL/vba/PersonalEventBuilder/obj_PEB_WordTemplates.cls.vba").read_text(encoding="utf-8-sig")
    assert "Word XML placeholder has no declared context provider" in parser
    assert "private_CompleteWordContext values, additionalInputs" in service
    assert service.index("    private_PrepareFoodSupport values") < service.index('    If exporter = "Word" Then')
    full = re.compile(settings["export.Word.duration.fullPattern"])
    operands = re.compile(settings["export.Word.duration.operandPattern"])
    for duration, total in [("1", 1), ("21", 21), ("15+2 (дорога) +1 (донація крові)", 18)]:
        match = full.fullmatch(duration)
        assert match and int(match.group(1)) + sum(int(item.group(1)) for item in operands.finditer(duration)) == total
    for duration in ["", "15+bad", "15-2"]:
        assert not full.fullmatch(duration)
    workbook_root = ROOT / "PROJECTS_EXCEL/workbook/PersonalEventBuilder"
    for group in ["context.lookups", "meta.lookups"]:
        for index in range(1, int(settings[f"export.Word.{group}.count"]) + 1):
            prefix = f"export.Word.{group}.{index}"
            book = openpyxl.load_workbook(workbook_root / settings[prefix + ".path"], data_only=True)
            sheet = book[settings[prefix + ".sheet"]]
            area = list(sheet[settings[prefix + ".range"]])
            headers = [cell.value for cell in area[0]]
            assert settings[prefix + ".keyHeader"] in headers and settings[prefix + ".valueHeader"] in headers
    print(f"PASS: all {len(root.findall('p:template', ns))} XML templates covered: {len(roots)} root placeholders, {len(row_keys)} row placeholders; context sources and duration grammar.")


def food_checks(tables, settings):
    offset = int(settings["export.Movement.foodSupport.minimumOrderOffsetDays"])
    assert offset >= 1
    def as_date(value):
        if isinstance(value, str):
            value = datetime.datetime.strptime(value, "%d.%m.%Y")
        if isinstance(value, (int, float)):
            value = openpyxl.utils.datetime.from_excel(value)
        return value.date() if isinstance(value, datetime.datetime) else value
    anchor = datetime.date(2026, 10, 15)
    for days, expected in [(-1, 1), (0, 1), (1, 1), (5, 5)]:
        assert max(anchor + datetime.timedelta(days=days), anchor + datetime.timedelta(days=offset)) == anchor + datetime.timedelta(days=max(expected, offset))
    steps = tables["tbTestSteps"]["rows"]
    checked = 0
    for row in tables["tbExpectedMovementRows"]["rows"]:
        for basis, food in [("Вибуття.Підстава", "Вибуття.Продовольче"), ("Прибуття.Підстава", "Прибуття.Продовольче")]:
            origins = [step for step in steps if row[basis] and f"№\u00a0{step['IncomingNumber']} від" in row[basis]]
            if not origins:
                continue
            assert len(origins) == 1
            step = origins[0]
            expected = max(as_date(step["EventDate"]), as_date(step["AcceptedOrderDate"]) + datetime.timedelta(days=offset))
            assert as_date(row[food]) == expected, f"Incorrect food support date: {row['StepId']}/{food}"
            checked += 1
    assert checked > 0
    for profile, index, header in [("Departure", 9, "Вибуття.Продовольче"), ("Arrival", 4, "Прибуття.Продовольче")]:
        assert settings[f"export.Movement.profiles.{profile}.fields.{index}.field"] == "FoodSupportDate"
        assert settings[f"export.Movement.profiles.{profile}.fields.{index}.header"] == header
    source = (ROOT / "PROJECTS_EXCEL/vba/PersonalEventBuilder/obj_PEB_ExportService.cls.vba").read_text(encoding="utf-8-sig")
    assert 'private_PrepareFoodSupport values' in source
    assert 'If eventDate > effectiveDate Then effectiveDate = eventDate' in source
    print(f"PASS: {checked} food support expectations, boundary dates and both export mappings.")


def order_checks(tables, settings):
    source = ROOT / "PROJECTS_EXCEL/workbook/PersonalEventBuilder" / settings["orders.source.path"]
    book = openpyxl.load_workbook(source, data_only=True)
    rows = list(book[settings["orders.source.sheet"]][settings["orders.source.range"]])
    headers = [cell.value for cell in rows[0]]
    date_column = headers.index(settings["orders.dateHeader"])
    number_column = headers.index(settings["orders.numberHeader"])
    pairs = set()
    dates = {}
    for row in rows[1:]:
        date = row[date_column].value
        number = row[number_column].value
        if date is None and number is None:
            continue
        assert str(number).isdigit(), "Invalid source order number"
        number = str(int(number))
        if isinstance(date, (int, float)):
            date = openpyxl.utils.datetime.from_excel(date, book.epoch)
        if isinstance(date, str):
            date = datetime.datetime.strptime(date, "%d.%m.%Y")
        assert isinstance(date, (datetime.datetime, datetime.date)), "Invalid source order date"
        date = date.date() if isinstance(date, datetime.datetime) else date
        key = (date.year, number)
        assert key not in dates or dates[key] == date, "Conflicting order dates in source"
        dates[key] = date
        pairs.add((number, date))
    for step in tables["tbTestSteps"]["rows"]:
        date = step["AcceptedOrderDate"]
        if isinstance(date, str):
            date = datetime.datetime.strptime(date, "%d.%m.%Y")
        if isinstance(date, (int, float)):
            date = openpyxl.utils.datetime.from_excel(date)
        date = date.date() if isinstance(date, datetime.datetime) else date
        assert (str(int(step["AcceptedOrderNumber"])), date) in pairs, f"Unknown order pair: {step['StepId']}"
    base = datetime.datetime.strptime(settings["test.fixtures.baseDate"], "%d.%m.%Y").date()
    assert ("115", base) not in pairs, "Negative fixture must remain absent"
    assert ("999999999", base) not in pairs
    text = (ROOT / "PROJECTS_EXCEL/vba/PersonalEventBuilder/obj_PEB_TestScenarios.cls.vba").read_text(encoding="utf-8-sig")
    prepare = text[text.index("Public Function TryPrepare("):text.index("End Function", text.index("Public Function TryPrepare("))]
    assert prepare.index("private_ResolveOrder updates") < prepare.index("private_SeedAbsence")
    assert 'm_orderResolver.TrySearch(query, year, matches, diagnostic)' in text
    assert 'If matches.Count = 0 Then Err.Raise' in text and 'If matches.Count <> 1 Then' in text
    print("PASS: all step order pairs exist in the UI source; shared resolver runs before history seeding.")


def directory_checks(tables, settings):
    workbook_root = ROOT / "PROJECTS_EXCEL/workbook/PersonalEventBuilder"
    movement = openpyxl.load_workbook(workbook_root / settings["export.Movement.path"])
    target = movement[settings["export.Movement.sheet"]].tables[settings["export.Movement.table"]]
    for column in target.tableColumns:
        formula = column.calculatedColumnFormula
        if formula is not None:
            assert "[2]!ТВ[" not in formula.attr_text and "[2]!Список[" not in formula.attr_text, f"External Movement self-reference: {column.name}"
    assert [column.name for column in target.tableColumns] == [header for header in tables["tbTestHistory"]["headers"] if header not in ("ScenarioId", "RowId", "TargetTable")]
    assert settings["environment.name"] == "TEST"
    test_root = (workbook_root / settings["test.scenarios.root"]).resolve()
    for key in ("export.Movement.path", "test.scenarios.path", "lookup.Absence.source.path"):
        assert (workbook_root / settings[key]).resolve().is_relative_to(test_root), f"Source outside TEST: {key}"
    assert (workbook_root / settings["export.Movement.path"]).resolve() != (workbook_root / settings["test.scenarios.productionPath"]).resolve()
    books = {}
    expected = unique(tables["tbExpectedWord"]["rows"], "StepId")
    for step in tables["tbTestSteps"]["rows"]:
        if step["Action"] != "PreviewWord":
            continue
        mode = runtime_mode(step, settings)
        for prefix in [f"export.Word.lookups.{index}" for index in range(1, int(settings["export.Word.lookups.count"]) + 1)] + [
            f"export.Word.{mode}.lookups.{index}"
            for index in range(1, int(settings[f"export.Word.{mode}.lookups.count"]) + 1)
        ]:
            input_field = settings[f"{prefix}.input"]
            if input_field.startswith("Report"):
                continue
            header = input_field
            lookup_path = (workbook_root / settings[f"{prefix}.path"]).resolve()
            if lookup_path not in books:
                books[lookup_path] = openpyxl.load_workbook(lookup_path, data_only=True)
            sheet = books[lookup_path][settings[f"{prefix}.sheet"]]
            left, top, right, bottom = range_boundaries(settings[f"{prefix}.range"])
            headers = [sheet.cell(top, column).value for column in range(left, right + 1)]
            key_column = left + headers.index(settings[f"{prefix}.keyHeader"])
            value_column = left + headers.index(settings[f"{prefix}.valueHeader"])
            matches = [sheet.cell(row, value_column).value for row in range(top + 1, min(bottom, sheet.max_row) + 1)
                       if sheet.cell(row, key_column).value == step[header]]
            if expected[step["StepId"]]["Outcome"] == "Success":
                assert len(matches) == 1 and matches[0], f"Missing preview lookup: {step['StepId']}/{prefix}"
            elif ".ToHospital.lookups" in prefix or ".FromHospital.lookups" in prefix:
                assert not matches, f"Negative hospital fixture unexpectedly resolves: {step['StepId']}"
    print("PASS: native Movement schema, TEST routing and all Word fixture directory dependencies.")


def registry_checks(book, tables, settings):
    kinds = settings["test.registry.kinds"].split("|")
    registry = tables[settings["test.registry.table"]]
    assert registry["sheet"] == settings["test.registry.sheet"]
    assert registry["headers"] == kinds + [settings["test.registry.runtimeHeader"]]
    templates = {node.get("name") for node in ET.parse(ROOT / "PROJECTS_EXCEL/workbook/PersonalEventBuilder/PersonalEventBuilderWordResultTemplates.xml").getroot()
                 if node.tag.endswith("}template")}
    for column, kind in enumerate(kinds, 1):
        expected = settings[f"test.registry.{kind}.values"].split("|")
        actual = [row[kind] for row in registry["rows"] if row[kind] is not None]
        assert actual == expected, f"Registry sheet is stale: {kind}"
        name = settings[f"test.registry.{kind}.rangeName"]
        assert len(name) <= 31 and name in book.defined_names
        sheet, address = next(book.defined_names[name].destinations)
        assert sheet == settings["test.registry.sheet"]
        left, top, right, bottom = range_boundaries(address)
        assert (left, top, right, bottom) == (column, 2, column, len(expected) + 1)
    for template in settings["test.registry.ModeId.values"].split("|"):
        assert template in templates, f"Unknown XML template: {template}"
        mode = settings[f"test.registry.ModeId.{template}.runtimeMode"]
        assert settings[f"export.Word.{mode}.templateName"] == template
    for binding in settings["test.registry.bindings"].split("|"):
        table_name, field, kind = binding.split(":")
        sheet = book[tables[table_name]["sheet"]]
        table = sheet.tables[table_name]
        left, top, right, bottom = range_boundaries(table.ref)
        column = left + tables[table_name]["headers"].index(field)
        cell = sheet.cell(top + 1, column).coordinate
        expected = settings[f"test.registry.{kind}.rangeName"]
        found = [rule for rule in sheet.data_validations.dataValidation if cell in rule.sqref]
        assert len(found) == 1 and found[0].type == "list" and found[0].formula1.lstrip("=") == expected, f"Missing dropdown: {binding}"
        assert found[0].errorStyle == "stop" and found[0].showErrorMessage
    loader = (ROOT / "PROJECTS_EXCEL/vba/PersonalEventBuilder/obj_PEB_TestScenarios.cls.vba").read_text(encoding="utf-8-sig")
    assert "private_SyncRegistry" in loader and 'fn_RegistryContains("ExpectedCode"' in loader
    print("PASS: TEST configuration registry, XML identifiers, named ranges and all dropdown bindings.")


def preservation_check(path):
    before = openpyxl.load_workbook(path, data_only=False)
    after = openpyxl.load_workbook(FIXTURE, data_only=False)
    for name in ("Тест_Історія",):
        assert list(before[name].values) == list(after[name].values), f"History changed: {name}"
    for name in ("Тест_Сценарії", "Тест_ЄЖОС"):
        omitted = {"Цільова таблиця"} if name == "Тест_Сценарії" else set()
        old_headers = [cell.value for cell in before[name][1]]
        included = [index for index, header in enumerate(old_headers) if header not in omitted]
        for old_row, new_row in zip(before[name].values, after[name].values):
            assert tuple(old_row[index] for index in included) == new_row, f"Original inputs changed: {name}"
    for name in ("Тест_Результати", "Тест_Записи", "Тест_ВхідніДані"):
        old_sheet = before[name]
        new_sheet = after[name]
        headers = [cell.value for cell in old_sheet[1]]
        scenario_column = headers.index("ScenarioId") + 1
        step_column = headers.index("StepId") + 1
        for row in old_sheet:
            for cell in row:
                expected = cell.value
                scenario = old_sheet.cell(cell.row, scenario_column).value
                if cell.column == step_column and scenario and isinstance(expected, int):
                    expected = f"{scenario}-M{expected:02}"
                assert new_sheet.cell(cell.row, cell.column).value == expected, f"Archive changed: {name}/{cell.coordinate}"
    print("PASS: original inputs, history and archive values preserved; historical StepId migrated.")


if __name__ == "__main__":
    settings = config(ROOT / "PROJECTS_EXCEL/config/PersonalEventBuilder/wsConfig.test.txt")
    production = config(ROOT / "PROJECTS_EXCEL/config/PersonalEventBuilder/wsConfig.txt")
    for key in settings:
        if key.startswith(("test.expected.", "test.scenarios.", "test.audit.")):
            assert settings[key] == production[key], f"PROD/TEST schema differs: {key}"
    book = openpyxl.load_workbook(FIXTURE)
    tables = load_tables(book)
    counts = validate(tables, settings)
    negative_checks(tables, settings)
    source_checks(settings)
    order_checks(tables, settings)
    food_checks(tables, settings)
    word_context_checks(settings)
    directory_checks(tables, settings)
    registry_checks(book, tables, settings)
    if len(sys.argv) > 1:
        preservation_check(Path(sys.argv[1]))
    print(f"PASS: {counts[0]} scenarios, {counts[1]} stable steps, {counts[2]} expected Movement rows, {counts[3]} Word expectations. Office was not started.")