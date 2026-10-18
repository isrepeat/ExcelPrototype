"""Read grouping scenarios from the common workbook for offline validation."""

from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.utils.cell import range_boundaries

from event_routes import WORKBOOK


def table_rows(book, name):
    matches = [(sheet, sheet.tables[name]) for sheet in book if name in sheet.tables]
    assert len(matches) == 1, "Missing or duplicate fixture table: " + name
    sheet, table = matches[0]
    left, top, right, bottom = range_boundaries(table.ref)
    headers = [sheet.cell(top, column).value for column in range(left, right + 1)]
    assert headers == [column.name for column in table.tableColumns]
    assert len({column.id for column in table.tableColumns}) == len(table.tableColumns)
    assert len(headers) == len(set(headers))
    rows = []
    for index in range(top + 1, bottom + 1):
        values = [sheet.cell(index, column).value for column in range(left, right + 1)]
        assert not any(sheet.cell(index, column).data_type == "e" for column in range(left, right + 1))
        if any(value is not None for value in values):
            rows.append(dict(zip(headers, values)))
    return sheet.title, headers, rows


def load_grouping_cases(settings):
    path = (WORKBOOK / settings["test.scenarios.path"]).resolve()
    book = openpyxl.load_workbook(path, data_only=False)
    _, _, all_cases = table_rows(book, settings["test.scenarios.cases.table"])
    identifiers = [row["ScenarioId"] for row in all_cases]
    assert len(identifiers) == len(set(identifiers))
    assert all(row["Suite"] in ("Basic", "Grouping") for row in all_cases)
    cases = {row["ScenarioId"]: row for row in all_cases if row["Suite"] == "Grouping"}
    tables = {}
    for kind in ("groupingSteps", "groupingExpected"):
        prefix = "test.scenarios." + kind
        sheet, headers, rows = table_rows(book, settings[prefix + ".table"])
        assert sheet == settings[prefix + ".sheet"]
        assert headers == settings[prefix + ".fields"].split("|")
        assert all(row["ScenarioId"] in cases for row in rows)
        tables[kind] = rows
    fields = settings["export.Word.grouping.stateFields"].split("|")
    step_ids = [row["StepId"] for row in tables["groupingSteps"]]
    assert len(step_ids) == len(set(step_ids))
    _, _, basic_steps = table_rows(book, settings["test.scenarios.steps.table"])
    assert set(step_ids).isdisjoint(row["StepId"] for row in basic_steps)
    root = ET.Element("scenarios")
    defaults = ET.SubElement(root, "defaults")
    for field in fields:
        ET.SubElement(defaults, "field", name=field).text = ""
    for identifier, case in cases.items():
        steps = sorted((row for row in tables["groupingSteps"] if row["ScenarioId"] == identifier), key=lambda row: row["StepOrder"])
        expected = [row for row in tables["groupingExpected"] if row["ScenarioId"] == identifier]
        assert steps and steps[0]["Action"] == "GroupAdd"
        assert [row["StepOrder"] for row in steps] == list(range(1, len(steps) + 1))
        assert expected and len({row["ExpectedEvents"] for row in expected}) == 1
        events = expected[0]["ExpectedEvents"]
        assert isinstance(events, int) and events > 0
        scenario = ET.SubElement(root, "scenario", id=identifier, events=str(events), name=case["Назва"])
        for row in steps:
            assert row["Action"] in ("GroupAdd", "GroupEdit")
            node = ET.SubElement(scenario, "add" if row["Action"] == "GroupAdd" else "edit", id=row["StepId"])
            if row["ExpectedError"]:
                node.set("error", row["ExpectedError"])
            if row["Action"] == "GroupAdd":
                for field in fields:
                    ET.SubElement(node, "field", name=field).text = str(row[field]) if row[field] is not None else ""
        for row in expected:
            assert row["Text"] and isinstance(row["Count"], int) and row["Count"] >= 0
            ET.SubElement(scenario, "expect", count=str(row["Count"])).text = row["Text"]
    book.close()
    return root