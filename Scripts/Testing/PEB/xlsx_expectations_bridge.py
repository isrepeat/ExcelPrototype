"""Чтение фикстур и сохранение неизменённых частей OOXML после редактирования."""

import copy
import datetime
import json
import posixpath
import os
import re
import sys
import zipfile
from lxml import etree as ET

import openpyxl
from openpyxl.utils.cell import get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.datetime import to_excel

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"
DOCREL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
NS = {"m": MAIN, "r": REL}


def dump(path):
    book = openpyxl.load_workbook(path, data_only=False)
    result = {}
    for sheet in book:
        rows = []
        for row in sheet.values:
            rows.append([
                to_excel(value, book.epoch) if isinstance(value, (datetime.datetime, datetime.date)) else value
                for value in row
            ])
        result[sheet.title] = {
            "rows": rows,
            "tables": [{"name": table.name, "ref": table.ref, "style": table.tableStyleInfo.name}
                       for table in sheet.tables.values()],
        }
    print(json.dumps(result, ensure_ascii=False))


def xml(data, name):
    return ET.fromstring(data[name])


def serialize(node):
    return ET.tostring(node, encoding="utf-8", xml_declaration=True)


def rel_path(part):
    folder, name = posixpath.split(part)
    return f"{folder}/_rels/{name}.rels"


def target(part, value):
    return value.lstrip("/") if value.startswith("/") else posixpath.normpath(posixpath.join(posixpath.dirname(part), value))


def sheet_parts(data):
    relationships = {item.get("Id"): target("xl/workbook.xml", item.get("Target"))
                     for item in xml(data, "xl/_rels/workbook.xml.rels")}
    return {item.get("name"): relationships[item.get(f"{{{DOCREL}}}id")]
            for item in xml(data, "xl/workbook.xml").find("m:sheets", NS)}


def merge_styles(original, generated):
    base = xml(original, "xl/styles.xml")
    added = xml(generated, "xl/styles.xml")
    offsets = {}
    for tag in ("fonts", "fills", "borders", "cellStyleXfs", "cellXfs", "dxfs"):
        node = base.find(f"m:{tag}", NS)
        if node is None:
            node = ET.SubElement(base, f"{{{MAIN}}}{tag}")
        offsets[tag] = len(node)
    formats = base.find("m:numFmts", NS)
    if formats is None:
        formats = ET.Element(f"{{{MAIN}}}numFmts")
        base.insert(0, formats)
    existing_formats = {item.get("formatCode"): int(item.get("numFmtId")) for item in formats}
    next_format = max([163] + list(existing_formats.values())) + 1
    format_ids = {}
    for item in added.findall("m:numFmts/m:numFmt", NS):
        code = item.get("formatCode")
        old_id = int(item.get("numFmtId"))
        if code not in existing_formats:
            clone = copy.deepcopy(item)
            clone.set("numFmtId", str(next_format))
            existing_formats[code] = next_format
            formats.append(clone)
            next_format += 1
        format_ids[old_id] = existing_formats[code]
    formats.set("count", str(len(formats)))
    for tag in offsets:
        node = base.find(f"m:{tag}", NS)
        for item in added.findall(f"m:{tag}/*", NS):
            clone = copy.deepcopy(item)
            if tag in ("cellStyleXfs", "cellXfs"):
                for attribute, collection in (("fontId", "fonts"), ("fillId", "fills"), ("borderId", "borders")):
                    if clone.get(attribute) is not None:
                        clone.set(attribute, str(int(clone.get(attribute)) + offsets[collection]))
                if clone.get("xfId") is not None:
                    clone.set("xfId", str(int(clone.get("xfId")) + offsets["cellStyleXfs"]))
                if clone.get("numFmtId") is not None:
                    old_id = int(clone.get("numFmtId"))
                    clone.set("numFmtId", str(format_ids.get(old_id, old_id)))
            node.append(clone)
        node.set("count", str(len(node)))
    original["xl/styles.xml"] = serialize(base)
    return offsets


def column_index(address):
    index = 0
    for character in address:
        if not character.isalpha():
            break
        index = index * 26 + ord(character) - ord("A") + 1
    return index


def remove_fixture_fields(data):
    removed = []
    fields = {"tbTestScenarios": "Цільова таблиця", "tbTestSteps": "PersonPosition"}
    for sheet_name, part in sheet_parts(data).items():
        relationships_part = rel_path(part)
        if relationships_part not in data:
            continue
        for relationship in xml(data, relationships_part):
            if not relationship.get("Type").endswith("/table"):
                continue
            table_part = target(part, relationship.get("Target"))
            table = xml(data, table_part)
            field = fields.get(table.get("name"))
            if not field:
                continue
            columns = table.find("m:tableColumns", NS)
            match = next((item for item in columns if item.get("name") == field), None)
            if match is None:
                continue
            index = column_index(table.get("ref").split(":")[0]) + list(columns).index(match)
            sheet = xml(data, part)
            if sheet.findall(".//m:f", NS):
                raise ValueError(f"Formula references require explicit migration: {sheet_name}")

            def shift(address):
                def replace(found):
                    number = column_index(found.group(1))
                    if number >= index:
                        number = max(1, number - 1)
                    return get_column_letter(number) + found.group(2)
                return re.sub(r"([A-Z]+)([0-9]+)", replace, address)

            for row in sheet.findall("m:sheetData/m:row", NS):
                for cell in list(row):
                    if column_index(cell.get("r")) == index:
                        row.remove(cell)
                    else:
                        cell.set("r", shift(cell.get("r")))
            for column in sheet.findall("m:cols/m:col", NS):
                start = int(column.get("min"))
                end = int(column.get("max"))
                if start == end == index:
                    column.getparent().remove(column)
                else:
                    column.set("min", str(start - (start > index)))
                    column.set("max", str(end - (end >= index)))
            for node in sheet.iter():
                for attribute in ("ref", "sqref", "activeCell", "topLeftCell"):
                    if node.get(attribute):
                        node.set(attribute, shift(node.get(attribute)))
            columns.remove(match)
            columns.set("count", str(len(columns)))
            for node in table.iter():
                if node.get("ref"):
                    node.set("ref", shift(node.get("ref")))
            data[part] = serialize(sheet)
            data[table_part] = serialize(table)
            removed.append(f"{sheet_name}/{field}")
    return removed


def clean_fixture(path):
    with zipfile.ZipFile(path) as archive:
        data = {name: archive.read(name) for name in archive.namelist()}
    removed = remove_fixture_fields(data)
    if not removed:
        print("No obsolete fixture fields found.")
        return
    temporary = str(path) + ".cleanup.tmp"
    try:
        with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, content in data.items():
                archive.writestr(name, content)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.remove(temporary)
    print("Removed fixture columns: " + ", ".join(removed))


def finalize_registry(path, config_path):
    with open(config_path, encoding="utf-8-sig") as source:
        settings = {parts[1].split("::", 1)[-1]: parts[2] for line in source.read().splitlines()
                    if line.strip() for parts in [line.split("\t")]}
    with zipfile.ZipFile(path) as archive:
        data = {name: archive.read(name) for name in archive.namelist()}
    tables = {}
    parts = sheet_parts(data)
    for name, sheet_part in parts.items():
        if rel_path(sheet_part) not in data:
            continue
        for relationship in xml(data, rel_path(sheet_part)):
            if relationship.get("Type").endswith("/table"):
                table_part = target(sheet_part, relationship.get("Target"))
                table = xml(data, table_part)
                tables[table.get("name")] = (sheet_part, table)
    registry_sheet = settings["test.registry.sheet"]
    kinds = settings["test.registry.kinds"].split("|")
    book = xml(data, "xl/workbook.xml")
    names = book.find("m:definedNames", NS)
    if names is None:
        names = ET.Element(f"{{{MAIN}}}definedNames")
        anchor = book.find("m:calcPr", NS)
        book.insert(list(book).index(anchor) if anchor is not None else len(book), names)
    for column, kind in enumerate(kinds, 1):
        name = settings[f"test.registry.{kind}.rangeName"]
        for old in list(names):
            if old.get("name") == name:
                names.remove(old)
        count = len(settings[f"test.registry.{kind}.values"].split("|"))
        node = ET.SubElement(names, f"{{{MAIN}}}definedName", {"name": name})
        node.text = f"'{registry_sheet}'!${get_column_letter(column)}$2:${get_column_letter(column)}${count + 1}"
    sheets = {}
    strings = xml(data, "xl/sharedStrings.xml").findall("m:si", NS) if "xl/sharedStrings.xml" in data else []
    step_part, step_table = tables[settings["test.scenarios.steps.table"]]
    steps = sheets.setdefault(step_part, xml(data, step_part))
    columns = [column.get("name") for column in step_table.find("m:tableColumns", NS)]
    left, top, right, bottom = range_boundaries(step_table.get("ref"))
    mode_column = left + columns.index(settings["test.scenarios.headers.mode"])
    modes = settings["test.registry.ModeId.values"].split("|")
    reverse = {settings[f"test.registry.ModeId.{mode}.runtimeMode"]: mode for mode in modes}
    for cell in steps.findall("m:sheetData/m:row/m:c", NS):
        if column_index(cell.get("r")) != mode_column or int(re.search(r"\d+", cell.get("r")).group()) <= top:
            continue
        if cell.get("t") == "s":
            value = "".join(strings[int(cell.find("m:v", NS).text)].itertext())
        elif cell.get("t") == "inlineStr":
            value = "".join(cell.find("m:is", NS).itertext())
        else:
            value = cell.findtext("m:v", default="", namespaces=NS)
        if not value:
            continue
        if value not in modes:
            if value not in reverse:
                raise ValueError(f"Cannot migrate test mode: {value}")
            for child in list(cell):
                cell.remove(child)
            cell.set("t", "inlineStr")
            string = ET.SubElement(cell, f"{{{MAIN}}}is")
            ET.SubElement(string, f"{{{MAIN}}}t").text = reverse[value]
    following = {"hyperlinks", "printOptions", "pageMargins", "pageSetup", "headerFooter", "rowBreaks", "colBreaks",
                 "customProperties", "cellWatches", "ignoredErrors", "smartTags", "drawing", "legacyDrawing",
                 "legacyDrawingHF", "picture", "oleObjects", "controls", "webPublishItems", "tableParts", "extLst"}
    for binding in settings["test.registry.bindings"].split("|"):
        table_name, field, kind = binding.split(":")
        part, table = tables[table_name]
        sheet = sheets.setdefault(part, xml(data, part))
        headers = [column.get("name") for column in table.find("m:tableColumns", NS)]
        left, top, right, bottom = range_boundaries(table.get("ref"))
        column = left + headers.index(field)
        validations = sheet.find("m:dataValidations", NS)
        if validations is None:
            validations = ET.Element(f"{{{MAIN}}}dataValidations")
            position = next((index for index, item in enumerate(sheet) if ET.QName(item).localname in following), len(sheet))
            sheet.insert(position, validations)
        for old in list(validations):
            for address in old.get("sqref").split():
                c1, r1, c2, r2 = range_boundaries(address)
                if c1 <= column <= c2 and r1 <= bottom and r2 >= top + 1:
                    validations.remove(old)
                    break
        node = ET.SubElement(validations, f"{{{MAIN}}}dataValidation", {
            "type": "list", "allowBlank": "1", "showDropDown": "0", "showErrorMessage": "1", "errorStyle": "stop",
            "errorTitle": settings["test.registry.errorTitle"], "error": settings["test.registry.errorMessage"],
            "sqref": f"{get_column_letter(column)}{top + 1}:{get_column_letter(column)}{bottom}"})
        ET.SubElement(node, f"{{{MAIN}}}formula1").text = "=" + settings[f"test.registry.{kind}.rangeName"]
        validations.set("count", str(len(validations)))
    data["xl/workbook.xml"] = serialize(book)
    for part, sheet in sheets.items():
        data[part] = serialize(sheet)
    temporary = str(path) + ".registry.tmp"
    try:
        with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, content in data.items():
                archive.writestr(name, content)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.remove(temporary)
    print(f"Registry named ranges and {len(settings['test.registry.bindings'].split('|'))} dropdown bindings updated.")


def finalize_movement_enums(path, expected_table):
    book_values = openpyxl.load_workbook(path, data_only=False)
    sheet_values = next(sheet for sheet in book_values if expected_table in sheet.tables)
    expected = sheet_values.tables[expected_table]
    enums = sheet_values.tables["tbMovementEnums"]
    left, top, right, bottom = range_boundaries(enums.ref)
    headers = [sheet_values.cell(top, column).value for column in range(left, right + 1)]
    event_column = left + headers.index("Подія")
    flag_column = left + headers.index("За межі")
    lookup = {}
    for index in range(top + 1, bottom + 1):
        event = sheet_values.cell(index, event_column).value
        if event is None or event in lookup:
            raise ValueError("Movement enumeration events must be populated and unique.")
        lookup[event] = sheet_values.cell(index, flag_column).value
    with zipfile.ZipFile(path) as archive:
        data = {name: archive.read(name) for name in archive.namelist()}
    part = sheet_parts(data)[sheet_values.title]
    sheet = xml(data, part)
    left, top, right, bottom = range_boundaries(expected.ref)
    headers = [column.name for column in expected.tableColumns]
    event_column = left + headers.index("Подія")
    destination_column = left + headers.index("Куди")
    flag_column = left + headers.index("За межі")
    current_event = f'{expected_table}[[#This Row],[Подія]]'
    current_destination = f'{expected_table}[[#This Row],[Куди]]'
    formula = f'IF({current_event}="","",IF(AND(OR({current_event}="Стаціонарне лікування",{current_event}="Відпустка для лікування"),{current_destination}="Медична рота А7383"),"Ні",_xlfn.XLOOKUP({current_event},tbMovementEnums[Подія],tbMovementEnums[За межі])))'
    for row in sheet.findall("m:sheetData/m:row", NS):
        index = int(row.get("r"))
        if not top < index <= bottom:
            continue
        event = sheet_values.cell(index, event_column).value
        destination = sheet_values.cell(index, destination_column).value
        if event and event not in lookup:
            raise ValueError(f"Unknown expected Movement event: {event}")
        value = "" if not event else ("Ні" if event in ("Стаціонарне лікування", "Відпустка для лікування") and destination == "Медична рота А7383" else lookup[event])
        if value is None:
            raise ValueError(f"Missing За межі for event: {event}")
        address = f"{get_column_letter(flag_column)}{index}"
        cell = next(cell for cell in row if cell.get("r") == address)
        for child in list(cell):
            if ET.QName(child).localname in ("f", "v", "is"):
                cell.remove(child)
        cell.set("t", "str")
        ET.SubElement(cell, f"{{{MAIN}}}f").text = formula
        ET.SubElement(cell, f"{{{MAIN}}}v").text = str(value)
    for relationship in xml(data, rel_path(part)):
        if relationship.get("Type").endswith("/table"):
            table_part = target(part, relationship.get("Target"))
            table = xml(data, table_part)
            if table.get("name") == expected_table:
                column = next(column for column in table.find("m:tableColumns", NS) if column.get("name") == "За межі")
                previous = column.find("m:calculatedColumnFormula", NS)
                if previous is not None:
                    column.remove(previous)
                ET.SubElement(column, f"{{{MAIN}}}calculatedColumnFormula").text = formula
                data[table_part] = serialize(table)
    book = xml(data, "xl/workbook.xml")
    names = book.find("m:definedNames", NS)
    if names is None:
        names = ET.Element(f"{{{MAIN}}}definedNames")
        anchor = book.find("m:calcPr", NS)
        book.insert(list(book).index(anchor) if anchor is not None else len(book), names)
    for name in list(names):
        if name.get("name") == "_PEB_MovementEvents":
            names.remove(name)
    ET.SubElement(names, f"{{{MAIN}}}definedName", {"name": "_PEB_MovementEvents"}).text = "tbMovementEnums[Подія]"
    validations = sheet.find("m:dataValidations", NS)
    if validations is None:
        validations = ET.Element(f"{{{MAIN}}}dataValidations")
        anchor = sheet.find("m:tableParts", NS)
        sheet.insert(list(sheet).index(anchor) if anchor is not None else len(sheet), validations)
    reference = f"{get_column_letter(event_column)}{top + 1}:{get_column_letter(event_column)}{bottom}"
    for validation in list(validations):
        if validation.get("sqref") == reference:
            validations.remove(validation)
    validation = ET.SubElement(validations, f"{{{MAIN}}}dataValidation", {
        "type": "list", "allowBlank": "1", "showErrorMessage": "1", "errorStyle": "stop", "sqref": reference})
    ET.SubElement(validation, f"{{{MAIN}}}formula1").text = "_PEB_MovementEvents"
    validations.set("count", str(len(validations)))
    data[part] = serialize(sheet)
    data["xl/workbook.xml"] = serialize(book)
    temporary = str(path) + ".enums.tmp"
    try:
        with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, content in data.items():
                archive.writestr(name, content)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.remove(temporary)
    print("Movement event dropdown and calculated За межі column updated.")


def format_food_dates(path, table_name):
    # Меняем только числовой формат двух колонок, сохраняя остальные свойства стилей.
    with zipfile.ZipFile(path) as archive:
        data = {name: archive.read(name) for name in archive.namelist()}
    styles = xml(data, "xl/styles.xml")
    formats = styles.find("m:numFmts", NS)
    if formats is None:
        formats = ET.Element(f"{{{MAIN}}}numFmts")
        styles.insert(0, formats)
    date_format = "dd.mm.yyyy"
    number_format = next((item.get("numFmtId") for item in formats if item.get("formatCode") == date_format), None)
    if number_format is None:
        number_format = str(max([163] + [int(item.get("numFmtId")) for item in formats]) + 1)
        ET.SubElement(formats, f"{{{MAIN}}}numFmt", {"numFmtId": number_format, "formatCode": date_format})
    formats.set("count", str(len(formats)))
    cell_styles = styles.find("m:cellXfs", NS)
    mapped_styles = {}
    found = False
    for part in sheet_parts(data).values():
        if rel_path(part) not in data:
            continue
        for relationship in xml(data, rel_path(part)):
            if not relationship.get("Type").endswith("/table"):
                continue
            table = xml(data, target(part, relationship.get("Target")))
            if table.get("name") != table_name:
                continue
            found = True
            left, top, _, bottom = range_boundaries(table.get("ref"))
            headers = [column.get("name") for column in table.find("m:tableColumns", NS)]
            columns = {left + headers.index(header) for header in ("Вибуття.Продовольче", "Прибуття.Продовольче")}
            sheet = xml(data, part)
            for row in sheet.findall("m:sheetData/m:row", NS):
                if not top < int(row.get("r")) <= bottom:
                    continue
                for cell in row:
                    if column_index(cell.get("r")) not in columns:
                        continue
                    index = int(cell.get("s", "0"))
                    if index not in mapped_styles:
                        clone = copy.deepcopy(cell_styles[index])
                        clone.set("numFmtId", number_format)
                        clone.set("applyNumberFormat", "1")
                        encoded = serialize(clone)
                        match = next((i for i, style in enumerate(cell_styles) if serialize(style) == encoded), None)
                        if match is None:
                            match = len(cell_styles)
                            cell_styles.append(clone)
                        mapped_styles[index] = str(match)
                    cell.set("s", mapped_styles[index])
            data[part] = serialize(sheet)
    if not found:
        raise ValueError(f"Missing expected Movement table: {table_name}")
    cell_styles.set("count", str(len(cell_styles)))
    data["xl/styles.xml"] = serialize(styles)
    temporary = str(path) + ".food.tmp"
    try:
        with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, content in data.items():
                archive.writestr(name, content)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.remove(temporary)


def localize_movement_formulas(path):
    # Удаляем ссылки копии тестовых формул на исходный РУХ, сохраняя остальные внешние связи.
    with zipfile.ZipFile(path) as archive:
        data = {name: archive.read(name) for name in archive.namelist()}
    changed = 0
    for name in list(data):
        if not (name.startswith("xl/worksheets/") or name.startswith("xl/tables/")) or not name.endswith(".xml"):
            continue
        node = xml(data, name)
        edited = False
        for formula in list(node.iter()):
            if ET.QName(formula).localname not in ("f", "calculatedColumnFormula") or not formula.text:
                continue
            text = formula.text
            for table in ("ТВ", "Список"):
                text = text.replace(f"[2]!{table}[", f"{table}[")
            if text != formula.text:
                formula.text = text
                if ET.QName(formula).localname == "f":
                    cell = formula.getparent()
                    cache = cell.find("m:v", NS)
                    if cache is not None:
                        cell.remove(cache)
                    cell.attrib.pop("t", None)
                changed += 1
                edited = True
        if edited:
            data[name] = serialize(node)
    if not changed:
        print("No external Movement self-references found.")
        return
    book = xml(data, "xl/workbook.xml")
    calculation = book.find("m:calcPr", NS)
    if calculation is None:
        calculation = ET.SubElement(book, f"{{{MAIN}}}calcPr")
    calculation.set("fullCalcOnLoad", "1")
    calculation.set("forceFullCalc", "1")
    data["xl/workbook.xml"] = serialize(book)
    temporary = str(path) + ".local.tmp"
    try:
        with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, content in data.items():
                archive.writestr(name, content)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.remove(temporary)
    print(f"Localized {changed} Movement formulas; styles and external link parts retained.")


def preserve_existing_cells(original, destination, edited, name):
    baseline = xml(original, destination)
    data = baseline.find("m:sheetData", NS)
    new_data = edited.find("m:sheetData", NS)
    old_rows = {int(row.get("r")): row for row in data}
    if name in ("Тест_Сценарії", "Тест_ЄЖОС"):
        last_row = max(old_rows)
        for row in new_data:
            if int(row.get("r")) > last_row:
                data.append(copy.deepcopy(row))
    else:
        step_column = 3 if name == "Тест_ВхідніДані" else 4
        for row in new_data:
            index = int(row.get("r"))
            for cell in row:
                column = column_index(cell.get("r"))
                if column != step_column and not (name == "Тест_Результати" and column > 22):
                    continue
                if index == 1 and column == step_column:
                    continue
                if index not in old_rows:
                    old_rows[index] = ET.SubElement(data, f"{{{MAIN}}}row", {"r": str(index)})
                target_row = old_rows[index]
                existing = next((item for item in target_row if item.get("r") == cell.get("r")), None)
                if existing is not None:
                    target_row.remove(existing)
                target_row.append(copy.deepcopy(cell))
        for row in data:
            row[:] = sorted(row, key=lambda cell: column_index(cell.get("r")))
    cells = data.findall("m:row/m:c", NS)
    last_column = max(column_index(cell.get("r")) for cell in cells)
    last_row = max(int(row.get("r")) for row in data)
    baseline.find("m:dimension", NS).set("ref", f"A1:{get_column_letter(last_column)}{last_row}")
    if name == "Тест_Результати":
        columns = baseline.find("m:cols", NS)
        for column in edited.findall("m:cols/m:col", NS):
            if int(column.get("min")) > 22:
                columns.append(copy.deepcopy(column))
    old_parts = baseline.find("m:tableParts", NS)
    old_parts[:] = [copy.deepcopy(item) for item in edited.find("m:tableParts", NS)]
    return baseline


def update_values_preserving_layout(original, generated, destination, source_part, strings):
    # Обновляем содержимое ячеек, сохраняя исходное оформление и объекты листа.
    if rel_path(destination) not in original or rel_path(source_part) not in generated:
        return False
    def tables(data, part):
        result = {}
        for relationship in xml(data, rel_path(part)):
            if relationship.get("Type").endswith("/table"):
                table_part = target(part, relationship.get("Target"))
                table = xml(data, table_part)
                result[table.get("name")] = (table_part, table)
        return result
    old_tables = tables(original, destination)
    new_tables = tables(generated, source_part)
    if old_tables.keys() != new_tables.keys():
        return False
    for name, (_, old) in old_tables.items():
        new = new_tables[name][1]
        if [column.get("name") for column in old.find("m:tableColumns", NS)] != [column.get("name") for column in new.find("m:tableColumns", NS)]:
            return False
        if range_boundaries(old.get("ref"))[:2] != range_boundaries(new.get("ref"))[:2]:
            return False
    baseline = xml(original, destination)
    edited = xml(generated, source_part)
    data = baseline.find("m:sheetData", NS)
    old_rows = {int(row.get("r")): row for row in data}
    styles = {}
    for cell in data.findall("m:row/m:c", NS):
        styles[column_index(cell.get("r"))] = cell.get("s")
    for row in edited.findall("m:sheetData/m:row", NS):
        index = int(row.get("r"))
        if index not in old_rows:
            preceding = [number for number in old_rows if number < index]
            previous = old_rows[max(preceding)] if preceding else None
            attributes = dict(previous.attrib) if previous is not None else {}
            attributes["r"] = str(index)
            attributes.pop("spans", None)
            old_rows[index] = ET.SubElement(data, f"{{{MAIN}}}row", attributes)
        target_row = old_rows[index]
        for cell in row:
            address = cell.get("r")
            existing = next((item for item in target_row if item.get("r") == address), None)
            if existing is not None and existing.find("m:f", NS) is not None:
                incoming = None
                if cell.get("t") == "s":
                    incoming = "".join(strings[int(cell.find("m:v", NS).text)].itertext())
                elif cell.get("t") == "inlineStr":
                    incoming = "".join(cell.find("m:is", NS).itertext())
                if incoming == "=" + existing.find("m:f", NS).text:
                    continue
            if existing is None:
                existing = ET.SubElement(target_row, f"{{{MAIN}}}c", {"r": address})
                style = styles.get(column_index(address))
                if style is not None:
                    existing.set("s", style)
            for child in list(existing):
                if ET.QName(child).localname in ("v", "f", "is"):
                    existing.remove(child)
            existing.attrib.pop("t", None)
            if cell.get("t") == "s":
                string = copy.deepcopy(strings[int(cell.find("m:v", NS).text)])
                string.tag = f"{{{MAIN}}}is"
                existing.append(string)
                existing.set("t", "inlineStr")
            else:
                if cell.get("t") is not None:
                    existing.set("t", cell.get("t"))
                for child in cell:
                    if ET.QName(child).localname in ("v", "f", "is"):
                        existing.append(copy.deepcopy(child))
        target_row[:] = sorted(target_row, key=lambda cell: column_index(cell.get("r")))
    data[:] = sorted(data, key=lambda row: int(row.get("r")))
    for name, (part, table) in old_tables.items():
        reference = new_tables[name][1].get("ref")
        left, _, right, old_bottom = range_boundaries(table.get("ref"))
        new_bottom = range_boundaries(reference)[3]
        for row in data:
            if new_bottom < int(row.get("r")) <= old_bottom:
                for cell in row:
                    if left <= column_index(cell.get("r")) <= right:
                        for child in list(cell):
                            if ET.QName(child).localname in ("v", "f", "is"):
                                cell.remove(child)
                        cell.attrib.pop("t", None)
        table.set("ref", reference)
        auto_filter = table.find("m:autoFilter", NS)
        if auto_filter is not None:
            auto_filter.set("ref", reference)
        original[part] = serialize(table)
    dimension = baseline.find("m:dimension", NS)
    cells = data.findall("m:row/m:c", NS)
    if dimension is not None and cells:
        dimension.set("ref", f"A1:{get_column_letter(max(column_index(cell.get('r')) for cell in cells))}{max(old_rows)}")
    original[destination] = serialize(baseline)
    return True


def assemble(original_path, generated_path, output_path, names, allow_layout_migration=False):
    with zipfile.ZipFile(original_path) as archive:
        original = {name: archive.read(name) for name in archive.namelist()}
    with zipfile.ZipFile(generated_path) as archive:
        generated = {name: archive.read(name) for name in archive.namelist()}
    baseline = dict(original)
    original_sheets = sheet_parts(original)
    generated_sheets = sheet_parts(generated)
    offsets = None
    strings = xml(generated, "xl/sharedStrings.xml").findall("m:si", NS) if "xl/sharedStrings.xml" in generated else []
    book = xml(original, "xl/workbook.xml")
    sheets = book.find("m:sheets", NS)
    relationships = xml(original, "xl/_rels/workbook.xml.rels")
    types = xml(original, "[Content_Types].xml")
    sheet_id = max(int(item.get("sheetId")) for item in sheets) + 1
    table_id = max(int(xml(original, name).get("id")) for name in original
                   if name.startswith("xl/tables/") and name.endswith(".xml")) + 1
    modified = {"xl/styles.xml", "xl/workbook.xml", "xl/_rels/workbook.xml.rels", "[Content_Types].xml"}
    insertion_index = list(sheets).index(next(item for item in sheets if item.get("name") == "Тест_Кроки")) + 1
    for name in names:
        source_part = generated_sheets[name]
        destination = original_sheets.get(name)
        if destination is not None and update_values_preserving_layout(original, generated, destination, source_part, strings):
            modified.add(destination)
            for relationship in xml(original, rel_path(destination)):
                if relationship.get("Type").endswith("/table"):
                    modified.add(target(destination, relationship.get("Target")))
            continue
        if destination is not None and not allow_layout_migration:
            raise ValueError(f"Sheet schema changed: {name}. Explicit --allow-layout-migration is required.")
        if offsets is None:
            offsets = merge_styles(original, generated)
        if destination is None:
            destination = f"xl/worksheets/sheet{sheet_id}.xml"
            relationship_id = f"rIdExpected{sheet_id}"
            sheet = ET.Element(f"{{{MAIN}}}sheet", {"name": name, "sheetId": str(sheet_id), f"{{{DOCREL}}}id": relationship_id})
            sheets.insert(insertion_index, sheet)
            insertion_index += 1
            ET.SubElement(relationships, f"{{{REL}}}Relationship", {
                "Id": relationship_id, "Type": f"{DOCREL}/worksheet", "Target": f"worksheets/sheet{sheet_id}.xml"})
            ET.SubElement(types, f"{{{CT}}}Override", {
                "PartName": f"/{destination}", "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"})
            sheet_id += 1
        node = xml(generated, source_part)
        for cell in node.findall(".//m:c", NS):
            if cell.get("s") is not None:
                cell.set("s", str(int(cell.get("s")) + offsets["cellXfs"]))
            if cell.get("t") == "s":
                value = cell.find("m:v", NS)
                string = copy.deepcopy(strings[int(value.text)])
                string.tag = f"{{{MAIN}}}is"
                cell.remove(value)
                cell.append(string)
                cell.set("t", "inlineStr")
        for rule in node.findall(".//m:cfRule", NS):
            if rule.get("dxfId") is not None:
                rule.set("dxfId", str(int(rule.get("dxfId")) + offsets["dxfs"]))
        if name in original_sheets and name in ("Тест_Сценарії", "Тест_ЄЖОС", "Тест_Результати", "Тест_Записи", "Тест_ВхідніДані"):
            node = preserve_existing_cells(original, destination, node, name)
        new_relationships = xml(generated, rel_path(source_part))
        original_tables = {}
        if rel_path(destination) in original:
            for relationship in xml(original, rel_path(destination)):
                if relationship.get("Type").endswith("/table"):
                    table_part = target(destination, relationship.get("Target"))
                    original_tables[xml(original, table_part).get("name")] = table_part
        for relationship in new_relationships:
            if not relationship.get("Type").endswith("/table"):
                raise ValueError(f"Unsupported relationship on edited sheet: {name}")
            table = xml(generated, target(source_part, relationship.get("Target")))
            table_part = original_tables.get(table.get("name"))
            if table_part:
                old_table = xml(original, table_part)
                table.set("id", old_table.get("id"))
                table.find("m:tableStyleInfo", NS).set("name", old_table.find("m:tableStyleInfo", NS).get("name"))
            else:
                table.find("m:tableStyleInfo", NS).set("name", "TableStyleDark_my")
                table_part = f"xl/tables/table{table_id}.xml"
                table.set("id", str(table_id))
                table_id += 1
                ET.SubElement(types, f"{{{CT}}}Override", {
                    "PartName": f"/{table_part}", "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.table+xml"})
            relationship.set("Target", posixpath.relpath(table_part, posixpath.dirname(destination)))
            original[table_part] = serialize(table)
            modified.add(table_part)
        original[destination] = serialize(node)
        original[rel_path(destination)] = serialize(new_relationships)
        modified.update((destination, rel_path(destination)))
    original["xl/workbook.xml"] = serialize(book)
    original["xl/_rels/workbook.xml.rels"] = serialize(relationships)
    original["[Content_Types].xml"] = serialize(types)
    for name, content in baseline.items():
        if name not in modified and original[name] != content:
            raise AssertionError(f"Untouched workbook part changed: {name}")
    remove_fixture_fields(original)
    with zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in original.items():
            archive.writestr(name, content)
    print(f"Preserved {len(baseline.keys() - modified)} untouched OOXML parts; updated {len(names)} sheets.")


if __name__ == "__main__":
    if sys.argv[1] == "dump":
        dump(sys.argv[2])
    elif sys.argv[1] == "assemble":
        assemble(*sys.argv[2:5], [name for name in sys.argv[5:] if name != "--allow-layout-migration"],
                 allow_layout_migration="--allow-layout-migration" in sys.argv[5:])
    elif sys.argv[1] == "clean-fixture":
        clean_fixture(sys.argv[2])
    elif sys.argv[1] == "finalize-registry":
        finalize_registry(sys.argv[2], sys.argv[3])
    elif sys.argv[1] == "finalize-movement-enums":
        finalize_movement_enums(sys.argv[2], sys.argv[3])
    elif sys.argv[1] == "format-food-dates":
        format_food_dates(sys.argv[2], sys.argv[3])
    elif sys.argv[1] == "localize-movement-formulas":
        localize_movement_formulas(sys.argv[2])
    else:
        raise SystemExit("Expected dump or assemble")