"""Resolve shared configuration for offline checks using the runtime import rules."""

from pathlib import Path
import re
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
WORKBOOK = ROOT / "PROJECTS_EXCEL/workbook/PersonalEventBuilder"


def rows(path):
    result = []
    for line in Path(path).read_text(encoding="utf-8-sig").splitlines():
        if not line:
            continue
        parts = line.split("\t")
        if len(parts) != 3:
            raise ValueError(f"Invalid configuration row: {path}: {line}")
        parts[0] = parts[0].strip()
        result.append(parts)
    return result


def expand(local, base, loader):
    entries = {}
    imported = set()
    for kind, key, value in local:
        kind = kind.strip()
        if kind.casefold() == "importcfg":
            prefix = "Import::workbooks."
            variable = key[len(prefix):]
            if not key.casefold().startswith(prefix.casefold()) or not re.fullmatch(r"__[A-Za-z0-9_]+", variable):
                raise ValueError(f"Invalid workbook import variable: {key}")
            path = (base / value).resolve()
            for shared_kind, shared_key, shared_value in loader(path):
                if shared_kind.casefold() == "importcfg":
                    raise ValueError("Nested imports are not allowed")
                if not shared_kind or not shared_key or shared_value in (None, ""):
                    raise ValueError("Empty shared configuration cell")
                shared_key = variable + "." + shared_key
                if shared_kind.casefold() == "ref" and not str(shared_value).startswith("__"):
                    shared_value = variable + "." + str(shared_value)
                if shared_key.casefold() in imported:
                    raise ValueError(f"Duplicate shared key: {shared_key}")
                imported.add(shared_key.casefold())
                entries[shared_key.casefold()] = shared_kind, str(shared_value), path.parent, shared_key
    local_keys = set()
    for kind, key, value in local:
        kind = kind.strip()
        if key.casefold() in local_keys:
            raise ValueError(f"Duplicate local key: {key}")
        local_keys.add(key.casefold())
        entries[key.casefold()] = kind, value, base, key
    resolved = {}
    visiting = set()

    def resolve(key):
        key = key.casefold()
        if key in resolved:
            return resolved[key]
        if key not in entries:
            raise ValueError(f"Referenced key not found: {key}")
        if key in visiting:
            raise ValueError(f"Reference cycle: {key}")
        visiting.add(key)
        kind, value, owner, _ = entries[key]
        if kind.casefold() == "ref":
            if value.casefold() in entries and entries[value.casefold()][0].casefold() == "importcfg":
                raise ValueError("Cannot reference an import directive")
            value = resolve(value)
            if not value:
                raise ValueError("Empty referenced configuration value")
        elif kind.casefold() in ("file", "importcfg"):
            if not value.strip():
                raise ValueError("Empty configuration file path")
            value = str((owner / value.replace('\\', '/')).resolve())
        visiting.remove(key)
        resolved[key] = value
        return value

    return {entry[3]: resolve(key) for key, entry in entries.items()}


def shared_rows(path):
    book = load_workbook(path, read_only=False, data_only=True)
    try:
        found = [(sheet, table) for sheet in book for table in sheet.tables.values() if table.name == "tbConfig"]
        if len(found) != 1:
            raise ValueError(f"Expected one tbConfig table: {path}")
        sheet, table = found[0]
        values = list(sheet[table.ref])
        headers = [cell.value for cell in values[0]]
        indexes = [headers.index(name) for name in ("Kind", "Key", "Value")]
        return [[row[index].value for index in indexes] for row in values[1:]]
    finally:
        book.close()


def load_config(path):
    return expand(rows(path), WORKBOOK, shared_rows)