"""Apply artifact-authored grouping fixture edits while preserving unrelated XLSX parts."""

import argparse
import hashlib
import json
from pathlib import Path
import re
import tempfile
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape
from zipfile import ZipFile

import openpyxl
from openpyxl.utils.cell import range_boundaries

from event_routes import ROOT, WORKBOOK, parse_config

NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def plan():
    settings = parse_config(ROOT / "PROJECTS_EXCEL/config/PersonalEventBuilder/wsConfig.test.txt")
    path = (WORKBOOK / settings["test.scenarios.path"]).resolve()
    book = openpyxl.load_workbook(path)
    changes = []
    for kind in ("groupingSteps", "groupingExpected"):
        sheet = book[settings[f"test.scenarios.{kind}.sheet"]]
        table = sheet.tables[settings[f"test.scenarios.{kind}.table"]]
        left, top, right, bottom = range_boundaries(table.ref)
        columns = {sheet.cell(top, col).value: col for col in range(left, right + 1)}
        for row in range(top + 1, bottom + 1):
            updates = {}
            if kind == "groupingSteps":
                if sheet.cell(row, columns["Action"]).value != "GroupAdd":
                    continue
                original = sheet.cell(row, columns["PersonId"]).value
                match = re.fullmatch(r"GROUP-P(\d{2})", original or "")
                if match:
                    person = f"ipn_{int(match[1]):04d}"
                    updates = {"PersonId": person, "IPN": person}
            else:
                case = sheet.cell(row, columns["ScenarioId"]).value
                text = sheet.cell(row, columns["Text"]).value
                if case == "G01-single" and text == "Відрядити лейтенанта":
                    updates = {"Text": "лейтенанта ТЕСТЕНКА Андрія"}
                if text == "Відрядити у" and case in ("G01-single", "G04-other-destination", "G05-other-departure-date"):
                    updates = {"Count": 1 if case == "G01-single" else 2}
            for field, value in updates.items():
                cell = sheet.cell(row, columns[field])
                assert cell.data_type != "f", "Fixture edit must not replace a formula"
                if cell.value != value:
                    changes.append({"sheet": sheet.title, "cell": cell.coordinate, "old": cell.value, "value": value})
    book.close()
    return {"path": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "changes": changes}


def apply(candidate, edits=None):
    if edits is None:
        edits = plan()
    path = Path(edits["path"])
    original = path.read_bytes()
    authored = openpyxl.load_workbook(candidate, data_only=False)
    for change in edits["changes"]:
        assert (authored[change["sheet"]][change["cell"]].value or "") == (change["value"] or ""), change
    authored.close()
    with ZipFile(path) as package:
        entries = package.infolist()
        payloads = {entry.filename: package.read(entry.filename) for entry in entries}
        comment = package.comment
    workbook = ET.fromstring(payloads["xl/workbook.xml"])
    rels = {node.get("Id"): node.get("Target") for node in ET.fromstring(payloads["xl/_rels/workbook.xml.rels"])}
    sheet_parts = {}
    for sheet in workbook.find("s:sheets", NS):
        target = rels[sheet.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")]
        sheet_parts[sheet.get("name")] = target.lstrip("/") if target.startswith("/") else "xl/" + target
    updated = {}
    for change in edits["changes"]:
        part = sheet_parts[change["sheet"]]
        xml = updated.get(part, payloads[part].decode("utf-8"))
        pattern = re.compile(r'<c\b(?=[^>]*\br="' + re.escape(change["cell"]) + r'")[^>]*>.*?</c>', re.S)
        matches = list(pattern.finditer(xml))
        assert len(matches) == 1, change
        old = matches[0].group()
        assert "<f" not in old
        attributes = re.sub(r'\s+t="[^"]*"', "", old[:old.index(">")])
        value = change["value"]
        if isinstance(value, str):
            new = attributes + ' t="inlineStr"><is><t xml:space="preserve">' + escape(value) + "</t></is></c>"
        else:
            new = attributes + "><v>" + str(value) + "</v></c>"
        updated[part] = xml[:matches[0].start()] + new + xml[matches[0].end():]
    for change in edits.get("tables", []):
        part = change["part"]
        original_xml = payloads[part].decode("utf-8")
        namespaces = re.findall(r'xmlns(?::([^=\s]+))?="([^"]+)"', original_xml)
        for prefix, uri in namespaces:
            ET.register_namespace(prefix or "", uri)
        table = ET.fromstring(payloads[part])
        table.set("ref", change["ref"])
        table.find("s:autoFilter", NS).set("ref", change["ref"])
        columns = table.find("s:tableColumns", NS)
        columns.clear()
        columns.set("count", str(len(change["headers"])))
        for index, name in enumerate(change["headers"], 1):
            ET.SubElement(columns, "{" + NS["s"] + "}tableColumn", id=str(index), name=name)
        xml = ET.tostring(table, encoding="unicode")
        for prefix, uri in namespaces:
            attribute = "xmlns" + (":" + prefix if prefix else "")
            if not re.search(r'\s' + re.escape(attribute) + r'="', xml):
                offset = xml.index(">")
                xml = xml[:offset] + f' {attribute}="{uri}"' + xml[offset:]
        updated[part] = xml
    backup = Path(tempfile.gettempdir()) / ("PEB-before-grouping-visuals-" + edits["sha256"][:16] + ".xlsx")
    if not backup.exists():
        backup.write_bytes(original)
    temporary = path.with_suffix(".visuals.tmp.xlsx")
    try:
        with ZipFile(temporary, "w") as package:
            package.comment = comment
            for entry in entries:
                package.writestr(entry, updated.get(entry.filename, "").encode("utf-8") if entry.filename in updated else payloads[entry.filename])
        with ZipFile(temporary) as package:
            assert package.testzip() is None
            assert all(package.read(name) == data for name, data in payloads.items() if name not in updated)
        checked = openpyxl.load_workbook(temporary)
        for change in edits["changes"]:
            assert (checked[change["sheet"]][change["cell"]].value or "") == (change["value"] or "")
        checked.close()
        assert path.read_bytes() == original, "Fixture changed during patch"
        temporary.replace(path)
    finally:
        if temporary.exists():
            temporary.unlink()
    print(f"PASS: {len(edits['changes'])} cells updated in {len(updated)} package parts; all other package parts preserved. Backup: {backup}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--candidate", type=Path)
    args = parser.parse_args()
    if args.candidate:
        apply(args.candidate)
    else:
        print(json.dumps(plan(), ensure_ascii=True))